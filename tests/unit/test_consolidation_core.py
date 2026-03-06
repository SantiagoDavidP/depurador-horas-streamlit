import io
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from backend.application.batch.batch_processor import BatchFileResult
from backend.application.consolidation.consolidator_integration import (
    generate_consolidated_from_batch_results,
    generate_individual_business_it_excel,
    validate_batch_results_for_consolidation,
)
from backend.application.consolidation.models import ConsolidatedReport
from backend.application.processing.models import ProcessorResult, ProcessorSummary
from backend.domain.models import ColumnMapping
from backend.infrastructure.config.collaborator_rates_repository import get_collaborator_rates_manager
from backend.infrastructure.reporting.consolidation.service import TimeSheetConsolidator
from backend.infrastructure.reporting.excel_timesheet_reporting_adapter import ExcelTimesheetReportingAdapter


def _tiny_png_bytes() -> bytes:
    # 1x1 transparent PNG
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc`\x00\x00"
        b"\x00\x02\x00\x01\xe2!\xbc3\x00\x00\x00\x00IEND\xaeB`\x82"
    )


def _build_df() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Fecha": ["2026-01-02", "2026-01-03"],
            "Proyecto": ["Proyecto X", "Proyecto X"],
            "Actividad": ["Desarrollo API", "Soporte QA y testing"],
            "Horas": [8.0, 2.5],
            "Tipo Hora": ["HN", "HE"],
            "Ticket": ["123", "ABC-1"],
            r"\NOVA\TecnologÃ­a - Documentos\IT\Desarrollo\DocumentaciÃ³n\2025": ["", None],
        }
    )


def _build_metadata(employee: str = "Luis Avila", company: str = "NOVA") -> dict:
    return {
        "employee": employee,
        "company": company,
        "month_name": "January",
        "year": 2026,
        "period_start": "2026-01-01",
        "period_end": "2026-01-31",
    }


def _build_processor_result(df: pd.DataFrame, metadata: dict) -> ProcessorResult:
    summary = ProcessorSummary(
        total_registros=len(df),
        total_errores=0,
        horas_totales=float(pd.to_numeric(df["Horas"], errors="coerce").fillna(0).sum()),
        dias_con_problemas_horas=0,
        errores_por_tipo={},
        errores_criticos=0,
        errores_advertencia=0,
        quality_score=100.0,
    )
    return ProcessorResult(
        workbook_bytes=b"dummy-xlsx",
        output_filename="salida.xlsx",
        validation_errors=[],
        corrections_log=[],
        summary=summary,
        corrected_dataframe=df.copy(),
        errors_dataframe=pd.DataFrame(columns=["fila", "tipo_error", "detalle"]),
        metadata=metadata.copy(),
    )


def test_generate_consolidated_report_creates_bytes_and_sheets(tmp_path):
    bit_logo = tmp_path / "bit.png"
    nova_logo = tmp_path / "nova.png"
    bit_logo.write_bytes(_tiny_png_bytes())
    nova_logo.write_bytes(_tiny_png_bytes())

    consolidator = TimeSheetConsolidator(
        cliente="NOVA - TI",
        collaborator_rates=get_collaborator_rates_manager(),
    )
    consolidator.logo_path = bit_logo
    consolidator.client_logo_path = nova_logo

    data = [
        (_build_df(), _build_metadata("Ana Perez", "NOVA - TI")),
        (_build_df(), _build_metadata("Bruno Diaz", "NOVA - TI")),
    ]

    report = consolidator.generate_consolidated_report(
        consultores_data=data,
        output_filename="Consolidado_Test.xlsx",
    )

    assert isinstance(report, ConsolidatedReport)
    assert report.output_filename == "Consolidado_Test.xlsx"
    assert report.consultores_incluidos == 2
    assert report.total_horas > 0
    assert len(report.workbook_bytes) > 100

    wb = load_workbook(io.BytesIO(report.workbook_bytes))
    assert "Resumen" in wb.sheetnames
    assert any("Ana Perez" in n or "Bruno Diaz" in n for n in wb.sheetnames)


def test_generate_single_consultant_report_for_baninter(tmp_path):
    bit_logo = tmp_path / "bit.png"
    ban_logo = tmp_path / "baninter.png"
    bit_logo.write_bytes(_tiny_png_bytes())
    ban_logo.write_bytes(_tiny_png_bytes())

    consolidator = TimeSheetConsolidator(
        cliente="BANINTER",
        collaborator_rates=get_collaborator_rates_manager(),
    )
    consolidator.logo_path = bit_logo
    consolidator.client_logo_path = ban_logo

    output = consolidator.generate_single_consultant_report(
        dataframe=_build_df(),
        metadata=_build_metadata("Pamela Chavez", "BANINTER"),
    )
    assert isinstance(output, (bytes, bytearray))
    assert len(output) > 100

    wb = load_workbook(io.BytesIO(output))
    assert wb.sheetnames
    assert "Pamela Chavez" in wb.sheetnames[0]


def test_consolidator_integration_and_validation_warnings():
    df_ok = _build_df()
    meta_ok = _build_metadata("Jonathan Chuquilla", "NOVA")
    ok_result = BatchFileResult(
        file_name="ok.xlsx",
        success=True,
        result=_build_processor_result(df_ok, meta_ok),
        sheet_name="Jonathan",
        client_id="cliente_bit",
        metadata=meta_ok,
    )
    failed_result = BatchFileResult(
        file_name="bad.xlsx",
        success=False,
        error="boom",
        client_id="cliente_bit",
    )

    is_valid, warnings = validate_batch_results_for_consolidation([ok_result, failed_result])
    assert is_valid is True
    assert any("fallaron" in w.lower() for w in warnings)

    report = generate_consolidated_from_batch_results(
        batch_results=[ok_result, failed_result],
        cliente="NOVA - TI",
        output_filename="Consolidado_Integration.xlsx",
        reporting_port=ExcelTimesheetReportingAdapter(),
        collaborator_rates=get_collaborator_rates_manager(),
    )
    assert report.output_filename == "Consolidado_Integration.xlsx"
    assert report.consultores_incluidos == 1
    assert len(report.workbook_bytes) > 100


def test_generate_individual_business_it_excel_from_batch_result():
    df_ok = _build_df()
    meta_ok = _build_metadata("Luis Avila", "BANINTER")
    processor_result = _build_processor_result(df_ok, meta_ok)
    # Simular columnas originales detectadas por el processor
    processor_result.date_column = "Fecha"  # type: ignore[attr-defined]
    processor_result.hours_column = "Horas"  # type: ignore[attr-defined]
    processor_result.description_column = "Actividad"  # type: ignore[attr-defined]

    batch_result = BatchFileResult(
        file_name="Resumen_Actividades.xlsx",
        success=True,
        result=processor_result,
        sheet_name="Luis Avila",
        client_id="cliente_talent",
        metadata=meta_ok,
    )

    data, filename = generate_individual_business_it_excel(
        batch_result,
        cliente="BANINTER",
        reporting_port=ExcelTimesheetReportingAdapter(),
        collaborator_rates=get_collaborator_rates_manager(),
    )
    assert isinstance(data, (bytes, bytearray))
    assert len(data) > 100
    assert filename.endswith(".xlsx")
    assert "BusinessIT" in filename
