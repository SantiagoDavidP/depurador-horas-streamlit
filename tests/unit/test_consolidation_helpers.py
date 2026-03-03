import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from backend.application.consolidation.service import TimeSheetConsolidator


def _tiny_png_bytes() -> bytes:
    return (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc`\x00\x00"
        b"\x00\x02\x00\x01\xe2!\xbc3\x00\x00\x00\x00IEND\xaeB`\x82"
    )


def _prepare_header(ws, title: str) -> None:
    ws.merge_cells("A1:N1")
    ws["A1"] = title
    fill = PatternFill(fill_type="solid", start_color="000000", end_color="000000")
    for r in (1, 2, 3):
        for c in range(1, 15):
            ws.cell(row=r, column=c).fill = fill
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 10


def test_logo_reference_and_normalization_helpers(tmp_path):
    bit_logo = tmp_path / "bit.png"
    nova_logo = tmp_path / "nova.png"
    bit_logo.write_bytes(_tiny_png_bytes())
    nova_logo.write_bytes(_tiny_png_bytes())

    consolidator = TimeSheetConsolidator(cliente="NOVA - TI")
    consolidator.logo_path = bit_logo
    consolidator.client_logo_path = nova_logo

    wb = Workbook()
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    _prepare_header(ws_resumen, "INFORME CONSOLIDADO DE ACTIVIDADES")
    consolidator._add_logo(ws_resumen, anchor="A1", logo_path=bit_logo, max_width=120, max_height=40)
    consolidator._add_logo(ws_resumen, anchor="N1", logo_path=nova_logo, max_width=120, max_height=40)

    refs = consolidator.extract_logo_reference_from_resumen(wb)
    assert refs is not None
    ref_bit, ref_nova = refs
    assert ref_bit["cx"] > 0
    assert ref_nova["cy"] > 0

    ws_individual = wb.create_sheet("1. Consultor")
    _prepare_header(ws_individual, "Informe de Actividades")
    consolidator._add_logo(ws_individual, anchor="A1", logo_path=bit_logo, max_width=100, max_height=30)
    consolidator._add_logo(ws_individual, anchor="N1", logo_path=nova_logo, max_width=100, max_height=30)

    consolidator.normalize_sheet_header_logos(
        ws_individual,
        ref_bit,
        ref_nova,
        bit_logo,
        nova_logo,
        title_texts=["Informe de Actividades", "Informe Consolidado de Actividades"],
    )
    assert len(getattr(ws_individual, "_images", [])) >= 2

    consolidator.normalize_header_logos(
        ws_individual,
        bit_logo,
        nova_logo,
        title_text="Informe de Actividades",
    )
    layout = consolidator.get_header_layout(ws_individual, title_text="Informe de Actividades")
    assert "leftLogoBox" in layout
    assert "rightLogoBox" in layout

    metrics = consolidator.get_header_metrics(ws_individual, title_text="Informe de Actividades")
    assert metrics["headerWidthPx"] > 0
    assert metrics["headerHeightPx"] > 0

    consolidator.normalize_header(
        ws_individual,
        bit_logo_path=bit_logo,
        nova_logo_path=nova_logo,
        title_text="Informe de Actividades",
    )
    assert len(getattr(ws_individual, "_images", [])) >= 2


def test_postprocess_and_cleanup_preserve_footer():
    consolidator = TimeSheetConsolidator(cliente="NOVA - TI")
    wb = Workbook()
    ws = wb.active
    ws.title = "Consultor 1"

    headers = ["Fecha", "Proyecto", "Horas", "Actividad", "Tipo Hora"]
    for i, header in enumerate(headers, start=1):
        ws.cell(row=9, column=i, value=header)

    # Filas de detalle (incluye una vacía que debería ocultarse)
    ws.cell(row=10, column=1, value="2026-01-02")
    ws.cell(row=10, column=3, value=8)
    ws.cell(row=10, column=4, value="Actividad 1")

    ws.cell(row=11, column=1, value=None)
    ws.cell(row=11, column=3, value=None)
    ws.cell(row=11, column=4, value=None)

    ws.cell(row=12, column=1, value="2026-01-03")
    ws.cell(row=12, column=3, value=2)
    ws.cell(row=12, column=4, value="Actividad 2")

    ws.cell(row=20, column=1, value="Elaborado por:")
    ws.cell(row=21, column=1, value="Nombre Aprobador")

    df = pd.DataFrame(
        {
            "Fecha": ["2026-01-02", "2026-01-03"],
            "Proyecto": ["P", "P"],
            "Horas": [8, 2],
            "Actividad": ["Actividad 1", "Actividad 2"],
            "Tipo Hora": ["HN", "HE"],
        }
    )

    consolidator.postprocess_hide_empty_table_rows_preserve_footer(wb, {ws.title: df})
    assert ws.row_dimensions[11].hidden is True
    assert ws.cell(row=20, column=1).value == "Elaborado por:"

    consolidator.clean_blank_rows_and_footer(ws)
    footer_values = [str(ws.cell(row=r, column=1).value or "") for r in range(1, ws.max_row + 1)]
    assert any(v.lower().startswith("elaborado por") for v in footer_values)
