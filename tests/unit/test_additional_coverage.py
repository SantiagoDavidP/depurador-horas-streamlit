import asyncio
import importlib
import os
from dataclasses import dataclass

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from backend.application.processing.debug_views import ProcessingDebugViewsMixin
from backend.application.processing.models import ProcessorSummary
from backend.application.processing.normalization import ProcessingNormalizationMixin, _safe_string
from backend.domain.models import ColumnMapping
from backend.infrastructure.auth import azure_ad_auth
from backend.infrastructure.config.collaborator_rates_repository import get_collaborator_rates_manager
from backend.infrastructure.export.excel_workbook_exporter import ExcelWorkbookExporter
from backend.infrastructure.health import health
from backend.infrastructure.reporting.consolidation.service import TimeSheetConsolidator


class _DummyProcessor(
    ProcessingNormalizationMixin,
    ProcessingDebugViewsMixin,
):
    pass


@dataclass
class _AzureADCfg:
    enabled: bool = False
    allowed_group_id: str | None = None
    client_id: str | None = None
    client_secret: str | None = None
    tenant_id: str | None = None
    redirect_uri: str = "http://localhost/callback"


@dataclass
class _AzureOpenAICfg:
    endpoint: str | None = None
    key: str | None = None
    api_version: str = "2024-05-01-preview"
    deployment: str = "gpt-4o-mini"


@dataclass
class _AzureStorageCfg:
    connection_string: str | None = None
    account_url: str | None = None
    credential: str | None = None
    container: str | None = "demo"


@dataclass
class _Settings:
    azure_ad: _AzureADCfg
    azure_openai: _AzureOpenAICfg
    azure_storage: _AzureStorageCfg


def test_protocol_modules_importable():
    # Cubre módulos Protocol con 0%: solo validar carga/import.
    m1 = importlib.import_module("api.application.ports.file_store_port")
    m2 = importlib.import_module("api.application.ports.user_auth_port")
    m3 = importlib.import_module("backend.application.ports.blob_storage_port")
    m4 = importlib.import_module("backend.application.ports.workbook_export_port")
    assert hasattr(m1, "FileStorePort")
    assert hasattr(m2, "UserAuthPort")
    assert hasattr(m3, "BlobStoragePort")
    assert hasattr(m4, "WorkbookExportPort")


def test_processing_mixins_normalization_debug_and_export():
    helper = _DummyProcessor()
    mapping = ColumnMapping(date="Fecha", hours="Horas", description="Actividad", project="Proyecto")

    df = pd.DataFrame(
        {
            "Fecha": ["2026-01-01", "", None, "2026-01-02"],
            "Horas": [8, 0, 2, 1],
            "Actividad": ["Trabajo", "total", "Soporte", ""],
            "Ticket": [" 123 ", "ABC 1", None, 45.0],
        }
    )
    filled = helper._fill_grouped_dates_for_payload(
        df,
        date_column="Fecha",
        hours_column="Horas",
        description_column="Actividad",
    )
    assert filled >= 1

    helper._normalize_ticket_columns(df)
    assert df["Ticket"].iloc[0] == "123"

    trimmed = helper._drop_empty_columns(pd.DataFrame({"A": [None, None], "B": ["", "  "], "C": [1, 2]}))
    assert list(trimmed.columns) == ["C"]

    raw_dbg = helper._group_hours_debug_df(df, mapping=mapping, row_numbers=[2, 3, 4, 5])
    merged_dbg = helper._merge_hours_debug(raw_dbg, raw_dbg, expected_hours=8.0)
    assert not merged_dbg.empty

    detail_dbg = helper._build_hours_debug_detail(
        df=df,
        mapping=mapping,
        row_numbers=[2, 3, 4, 5],
        validation_errors=[{"tipo_error": "horas_incorrectas", "fecha": "2026-01-01"}],
        stage_label="post-clean",
    )
    assert not detail_dbg.empty

    summary = ProcessorSummary(
        total_registros=4,
        total_errores=1,
        horas_totales=11.0,
        dias_con_problemas_horas=1,
        errores_por_tipo={"horas_incorrectas": 1},
        metadata_removidas=0,
        errores_criticos=1,
        errores_advertencia=0,
        quality_score=65.0,
        ai_summary={"diagnostico": "x", "acciones": ["a1", "a2"], "tiempo_estimado": "15 min"},
    )
    workbook_bytes = ExcelWorkbookExporter().export_workbook(
        df,
        pd.DataFrame([{"tipo_error": "horas_incorrectas"}]),
        summary,
    )
    assert len(workbook_bytes) > 100

    assert _safe_string(None) == ""


def test_postprocess_helpers_and_logo_anchor_paths(tmp_path):
    bit_logo = tmp_path / "bit.png"
    bit_logo.write_bytes(
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc`\x00\x00"
        b"\x00\x02\x00\x01\xe2!\xbc3\x00\x00\x00\x00IEND\xaeB`\x82"
    )

    c = TimeSheetConsolidator(
        cliente="NOVA - TI",
        collaborator_rates=get_collaborator_rates_manager(),
    )
    ws = Workbook().active
    ws.title = "Demo"
    ws.merge_cells("A1:H1")
    ws["A1"] = "Informe de Actividades"
    fill = PatternFill(fill_type="solid", start_color="000000", end_color="000000")
    for r in (1, 2, 3):
        for col in range(1, 9):
            ws.cell(row=r, column=col).fill = fill

    # _detect_header_band + _position_header_logo (_logo_anchor.py)
    band = c._detect_header_band(ws, ["informe de actividades"])
    assert band[0] == 1
    c._position_header_logo(
        ws,
        bit_logo,
        title_texts=["informe de actividades"],
        max_width=120,
        max_height=40,
    )
    assert len(getattr(ws, "_images", [])) >= 1

    footer_block = {
        "data_end_row": 10,
        "start_row": 12,
        "max_row": 13,
        "row_heights": {12: 20},
        "cells": [
            {
                "row": 12,
                "col": 1,
                "value": "Elaborado por:",
                "font": None,
                "fill": None,
                "border": None,
                "alignment": None,
                "number_format": None,
            }
        ],
        "merges": [(12, 1, 12, 3)],
    }
    fmt = c._detect_footer_format(footer_block)
    assert fmt in {"alfredo", "simple"}
    c._apply_footer_block(ws, footer_block, last_data_row=15)
    assert ws.cell(row=17, column=1).value == "Elaborado por:"

    c.consultores_metrics = [
        type("M", (), {"metadata": {"area": "CD"}})(),
        type("M", (), {"metadata": {"area": "TI"}})(),
    ]
    c._resolve_cliente_from_metrics()
    assert c.cliente in {"NOVA - Centro Digital", "BIT Nova - TI"}


def test_health_internal_openai_and_sql_checks(monkeypatch):
    checker = health.HealthChecker()

    # OpenAI not configured
    from config import settings as cfg_settings

    monkeypatch.setattr(
        cfg_settings,
        "get_settings",
        lambda: _Settings(_AzureADCfg(), _AzureOpenAICfg(endpoint=None, key=None), _AzureStorageCfg()),
    )
    openai_not_cfg = asyncio.run(checker._check_openai_connection())
    assert openai_not_cfg["healthy"] is True

    # OpenAI configured path
    import openai

    monkeypatch.setattr(
        cfg_settings,
        "get_settings",
        lambda: _Settings(_AzureADCfg(), _AzureOpenAICfg(endpoint="https://example", key="k"), _AzureStorageCfg()),
    )
    monkeypatch.setattr(openai, "AsyncAzureOpenAI", lambda **kwargs: object())
    openai_ok = asyncio.run(checker._check_openai_connection())
    assert openai_ok["healthy"] is True

    # SQL branch disabled by env
    old = os.environ.get("DATA_SOURCE_TYPE")
    os.environ["DATA_SOURCE_TYPE"] = "excel"
    sql_skip = asyncio.run(checker._check_sql_connection())
    assert sql_skip["healthy"] is True

    # SQL configured branch
    os.environ["DATA_SOURCE_TYPE"] = "sql"
    from backend.adapters import DataSourceConfig
    from backend.adapters import sql_adapter

    class _Adapter:
        def __init__(self, config: DataSourceConfig):
            self.server = "server"
            self.database = "db"

    monkeypatch.setattr(sql_adapter, "FabricSQLAdapter", _Adapter)
    sql_ok = asyncio.run(checker._check_sql_connection())
    assert sql_ok["healthy"] is True
    if old is None:
        os.environ.pop("DATA_SOURCE_TYPE", None)
    else:
        os.environ["DATA_SOURCE_TYPE"] = old


def test_health_not_ready_error_and_singleton_paths(monkeypatch):
    checker = health.HealthChecker()

    # Rama not_ready (todos los checks no saludables)
    monkeypatch.setattr(health.psutil, "virtual_memory", lambda: type("M", (), {"percent": 95.0, "available": 1 * 1024**3})())
    monkeypatch.setattr(health.psutil, "cpu_percent", lambda interval=0.1: 99.0)
    monkeypatch.setattr(
        checker,
        "_check_openai_connection",
        lambda: asyncio.sleep(0, result={"healthy": False, "status": "error"}),
    )
    monkeypatch.setattr(
        checker,
        "_check_sql_connection",
        lambda: asyncio.sleep(0, result={"healthy": False, "status": "error"}),
    )
    not_ready = asyncio.run(checker.readiness_check())
    assert not_ready["status"] == "not_ready"

    # Rama error (excepción en readiness)
    def _raise_memory():
        raise RuntimeError("boom")

    monkeypatch.setattr(health.psutil, "virtual_memory", _raise_memory)
    errored = asyncio.run(checker.readiness_check())
    assert errored["status"] == "error"

    # Singleton
    health._health_checker = None
    h1 = health.get_health_checker()
    h2 = health.get_health_checker()
    assert h1 is h2


def test_health_openai_sql_exception_and_render_paths(monkeypatch):
    checker = health.HealthChecker()

    from config import settings as cfg_settings
    import openai

    # OpenAI: forzar excepción para cubrir rama except
    monkeypatch.setattr(
        cfg_settings,
        "get_settings",
        lambda: _Settings(_AzureADCfg(), _AzureOpenAICfg(endpoint="https://x", key="k"), _AzureStorageCfg()),
    )
    monkeypatch.setattr(openai, "AsyncAzureOpenAI", lambda **kwargs: (_ for _ in ()).throw(RuntimeError("openai error")))
    openai_err = asyncio.run(checker._check_openai_connection())
    assert openai_err["healthy"] is False

    # SQL: rama misconfigured
    old = os.environ.get("DATA_SOURCE_TYPE")
    os.environ["DATA_SOURCE_TYPE"] = "sql"
    from backend.adapters import sql_adapter

    class _BadAdapter:
        def __init__(self, config):
            self.server = None
            self.database = None

    monkeypatch.setattr(sql_adapter, "FabricSQLAdapter", _BadAdapter)
    sql_bad = asyncio.run(checker._check_sql_connection())
    assert sql_bad["status"] == "misconfigured"

    # SQL: rama exception
    monkeypatch.setattr(sql_adapter, "FabricSQLAdapter", lambda config: (_ for _ in ()).throw(RuntimeError("sql boom")))
    sql_err = asyncio.run(checker._check_sql_connection())
    assert sql_err["healthy"] is False
    if old is None:
        os.environ.pop("DATA_SOURCE_TYPE", None)
    else:
        os.environ["DATA_SOURCE_TYPE"] = old

    # Render (si Streamlit está disponible)
    if hasattr(health, "render_health_status"):
        class _Ctx:
            def __enter__(self):
                return self

            def __exit__(self, exc_type, exc, tb):
                return False

        class _St:
            sidebar = _Ctx()

            @staticmethod
            def expander(*args, **kwargs):
                return _Ctx()

            @staticmethod
            def success(*args, **kwargs):
                return None

            @staticmethod
            def error(*args, **kwargs):
                return None

            @staticmethod
            def metric(*args, **kwargs):
                return None

            @staticmethod
            def caption(*args, **kwargs):
                return None

        monkeypatch.setattr(health, "st", _St)
        monkeypatch.setattr(health, "get_health_checker", lambda: checker)
        monkeypatch.setattr(
            checker,
            "liveness_check",
            lambda: {"status": "alive", "uptime_seconds": 10},
        )
        monkeypatch.setattr(health.psutil, "virtual_memory", lambda: type("M", (), {"percent": 50.0})())
        monkeypatch.setattr(health.psutil, "cpu_percent", lambda interval=0.1: 20.0)

        fn = getattr(health.render_health_status, "__wrapped__", health.render_health_status)
        fn()

        # Rama "dead" en render
        monkeypatch.setattr(
            checker,
            "liveness_check",
            lambda: {"status": "dead", "uptime_seconds": 0},
        )
        fn()


def test_health_liveness_exception_branch(monkeypatch):
    checker = health.HealthChecker()
    real_datetime = health.datetime

    class _FlakyDateTime:
        calls = 0

        @classmethod
        def utcnow(cls):
            cls.calls += 1
            if cls.calls == 1:
                raise RuntimeError("clock fail")
            return real_datetime.utcnow()

    monkeypatch.setattr(health, "datetime", _FlakyDateTime)
    result = checker.liveness_check()
    assert result["status"] == "dead"


def test_azure_ad_require_authentication_login_screen(monkeypatch):
    class _MsalApp:
        def get_authorization_request_url(self, scopes, redirect_uri):
            return "https://login.test/auth"

    class _Ctx:
        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, tb):
            return False

    class _Query(dict):
        def clear(self):
            super().clear()

    class _FakeSt:
        session_state = {}
        query_params = _Query()

        @staticmethod
        def error(*args, **kwargs):
            return None

        @staticmethod
        def info(*args, **kwargs):
            return None

        @staticmethod
        def markdown(*args, **kwargs):
            return None

        @staticmethod
        def caption(*args, **kwargs):
            return None

        @staticmethod
        def columns(*args, **kwargs):
            return [_Ctx(), _Ctx(), _Ctx()]

        @staticmethod
        def button(*args, **kwargs):
            return False

        @staticmethod
        def spinner(*args, **kwargs):
            return _Ctx()

    monkeypatch.setattr(azure_ad_auth, "st", _FakeSt)
    monkeypatch.setattr(
        azure_ad_auth,
        "get_settings",
        lambda: _Settings(
            _AzureADCfg(
                enabled=True,
                client_id="cid",
                client_secret="secret",
                tenant_id="tenant",
                redirect_uri="http://localhost/callback",
            ),
            _AzureOpenAICfg(),
            _AzureStorageCfg(),
        ),
    )
    monkeypatch.setattr(
        azure_ad_auth.msal,
        "ConfidentialClientApplication",
        lambda *args, **kwargs: _MsalApp(),
    )
    ok, user = azure_ad_auth.require_authentication()
    assert ok is False
    assert user is None
