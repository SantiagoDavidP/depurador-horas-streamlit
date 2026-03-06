from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ============================================================
# PALETA DE COLORES PROFESIONAL NOVA-TI
# ============================================================
COLOR_LIME = "8BC34A"              # Verde lima (titulos, acentos)
COLOR_DARK_GREEN = "689F38"        # Verde oscuro (totales, bordes)
COLOR_BLACK = "1A1A1A"             # Negro (fondos oscuros)
COLOR_SOFT_GREEN = "E8F5E9"        # Verde suave (filas alternadas)
COLOR_WHITE = "FFFFFF"             # Blanco
COLOR_LIGHT_GRAY = "F5F5F5"        # Gris claro (celdas de info)
COLOR_ORANGE = "F4A640"            # Naranja suave (encabezados resaltados)
COLOR_BORDER_LIGHT = "E0E0E0"      # Bordes sutiles grises
COLOR_CYAN = "1BA9D5"              # Azul/cian (totales destacados)
COLOR_RED = "E74C3C"               # Rojo (total general)
COLOR_SENIOR_ROW = "E6D9F2"        # Morado claro (Senior)
COLOR_SEMISENIOR_ROW = "D6ECFA"    # Azul claro (Semisenior)
COLOR_JUNIOR_ROW = "E8F5E9"        # Verde claro (Junior)

# ============================================================
# ESTILOS TIPOGRÁFICOS
# ============================================================
# Título principal del reporte
TITLE_FONT = Font(name="Calibri", size=22, bold=True, color=COLOR_LIME)

# Subtítulo (nombre del cliente)
SUBTITLE_FONT = Font(name="Calibri", size=12, bold=True, color=COLOR_WHITE)

# Etiquetas de información general
INFO_LABEL_FONT = Font(name="Calibri", size=11, bold=True, color=COLOR_BLACK)
INFO_VALUE_FONT = Font(name="Calibri", size=11, bold=False, color=COLOR_BLACK)

# Encabezados de tabla
TABLE_HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=COLOR_LIME)
TABLE_HEADER_FILL = PatternFill(start_color=COLOR_BLACK, end_color=COLOR_BLACK, fill_type="solid")
TABLE_HEADER_ORANGE_FONT = Font(name="Calibri", size=11, bold=True, color=COLOR_BLACK)
TABLE_HEADER_ORANGE_FILL = PatternFill(start_color=COLOR_ORANGE, end_color=COLOR_ORANGE, fill_type="solid")

# Filas de datos
DATA_FONT = Font(name="Calibri", size=10, bold=False, color=COLOR_BLACK)
ALTERNATE_ROW_FILL = PatternFill(start_color=COLOR_SOFT_GREEN, end_color=COLOR_SOFT_GREEN, fill_type="solid")

# Fila de totales
TOTAL_ROW_FILL = PatternFill(start_color=COLOR_DARK_GREEN, end_color=COLOR_DARK_GREEN, fill_type="solid")
TOTAL_ROW_FONT = Font(name="Calibri", size=11, bold=True, color=COLOR_WHITE)
TOTAL_ROW_HIGHLIGHT_FILL = PatternFill(start_color=COLOR_CYAN, end_color=COLOR_CYAN, fill_type="solid")
TOTAL_ROW_GRAND_FILL = PatternFill(start_color=COLOR_RED, end_color=COLOR_RED, fill_type="solid")
SENIOR_ROW_FILL = PatternFill(start_color=COLOR_SENIOR_ROW, end_color=COLOR_SENIOR_ROW, fill_type="solid")
SEMISENIOR_ROW_FILL = PatternFill(start_color=COLOR_SEMISENIOR_ROW, end_color=COLOR_SEMISENIOR_ROW, fill_type="solid")
JUNIOR_ROW_FILL = PatternFill(start_color=COLOR_JUNIOR_ROW, end_color=COLOR_JUNIOR_ROW, fill_type="solid")

# ============================================================
# ESTILOS DE BORDES
# ============================================================
# Bordes sutiles para tabla
THIN_BORDER = Border(
    left=Side(style="thin", color=COLOR_BORDER_LIGHT),
    right=Side(style="thin", color=COLOR_BORDER_LIGHT),
    top=Side(style="thin", color=COLOR_BORDER_LIGHT),
    bottom=Side(style="thin", color=COLOR_BORDER_LIGHT),
)

# Borde superior grueso para fila de totales
THICK_TOP_BORDER = Border(
    left=Side(style="thin", color=COLOR_BORDER_LIGHT),
    right=Side(style="thin", color=COLOR_BORDER_LIGHT),
    top=Side(style="thick", color=COLOR_LIME),
    bottom=Side(style="thin", color=COLOR_BORDER_LIGHT),
)

# Bordes para sección de información
INFO_BORDER = Border(
    left=Side(style="thin", color=COLOR_BORDER_LIGHT),
    right=Side(style="thin", color=COLOR_BORDER_LIGHT),
    top=Side(style="thin", color=COLOR_BORDER_LIGHT),
    bottom=Side(style="thin", color=COLOR_BORDER_LIGHT),
)

# ============================================================
# ESTILOS DE ALINEACIÓN
# ============================================================
CENTER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=False)
RIGHT_ALIGNMENT = Alignment(horizontal="right", vertical="center", wrap_text=False)

# Fondo para sección de información
INFO_FILL = PatternFill(start_color=COLOR_LIGHT_GRAY, end_color=COLOR_LIGHT_GRAY, fill_type="solid")
