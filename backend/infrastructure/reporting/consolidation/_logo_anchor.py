from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class LogoAnchorMixin:
    @staticmethod
    def _column_width_to_pixels(width: Optional[float]) -> int:
        # Aproximación estándar: Excel width -> px
        if width is None:
            width = 8.43
        return int(round(width * 7 + 5))

    @staticmethod
    def _row_height_to_pixels(height: Optional[float], default_height: Optional[float]) -> int:
        points = height if height is not None else (default_height or 15.0)
        return int(round(points * (96 / 72)))

    def _column_width_to_emu(self, width: Optional[float]) -> int:
        emu_per_px = 9525
        return self._column_width_to_pixels(width) * emu_per_px

    def _row_height_to_emu(self, height: Optional[float], default_height: Optional[float]) -> int:
        points = height if height is not None else (default_height or 15.0)
        return int(round(points * 12700))

    def _detect_header_band(
        self,
        ws: Worksheet,
        title_texts: List[str],
    ) -> Tuple[int, int, int, int]:
        """Detecta banda de header usando el rango merged del tÃ­tulo."""
        title_texts = [t.lower() for t in title_texts]
        title_cell = None
        for row in ws.iter_rows(min_row=1, max_row=8):
            for cell in row:
                if cell.value and any(t in str(cell.value).lower() for t in title_texts):
                    title_cell = cell
                    break
            if title_cell:
                break

        if not title_cell:
            return (1, 1, 1, ws.max_column or 1)

        # Buscar rango merged que contiene el tÃ­tulo
        merged_range = None
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= title_cell.row <= mr.max_row and mr.min_col <= title_cell.column <= mr.max_col:
                merged_range = mr
                break

        if not merged_range:
            return (title_cell.row, title_cell.row, 1, ws.max_column or 1)

        r1 = merged_range.min_row
        r2 = merged_range.max_row
        c1 = merged_range.min_col
        c2 = merged_range.max_col

        # Extender hacia abajo mientras haya fill en la banda
        max_scan = min(r1 + 6, ws.max_row)
        for r in range(r1, max_scan + 1):
            has_fill = False
            for c in range(c1, c2 + 1):
                cell = ws.cell(row=r, column=c)
                if cell.fill and cell.fill.fill_type:
                    has_fill = True
                    break
            if has_fill:
                r2 = r
            else:
                break
        return (r1, r2, c1, c2)

    def _position_header_logo(
        self,
        ws: Worksheet,
        img_path: Path,
        *,
        header_rows: Optional[Tuple[int, int]] = None,
        logo_block_cols: Optional[Tuple[int, int]] = None,
        title_texts: Optional[List[str]] = None,
        max_width: int = 140,
        max_height: int = 48,
    ) -> None:
        """Inserta logo centrado vertical y horizontalmente dentro del header."""
        if not img_path.exists():
            return

        # Detectar banda de header
        if header_rows is None or logo_block_cols is None:
            r1, r2, c1, c2 = self._detect_header_band(
                ws, title_texts or ["informe de actividades", "informe consolidado de actividades"]
            )
            if header_rows is None:
                header_rows = (r1, r2)
            if logo_block_cols is None:
                # por defecto usa las Ãºltimas 2 columnas del header
                logo_block_cols = (max(c2 - 1, c1), c2)

        r1, r2 = header_rows
        c1, c2 = logo_block_cols

        # Eliminar logos previos en esa zona
        kept_images = []
        for img in getattr(ws, "_images", []):
            anchor = img.anchor
            img_col = None
            img_row = None
            if isinstance(anchor, str):
                try:
                    col_letter, row = re.match(r"([A-Z]+)(\d+)", anchor).groups()
                    img_col = column_index_from_string(col_letter)
                    img_row = int(row)
                except Exception:
                    img_col = None
            elif hasattr(anchor, "_from"):
                img_col = anchor._from.col + 1
                img_row = anchor._from.row + 1

            if img_col is not None and img_row is not None:
                if c1 <= img_col <= c2 and r1 <= img_row <= r2:
                    continue
            kept_images.append(img)
        ws._images = kept_images

        img = Image(str(img_path))
        if not img.width or not img.height:
            return

        # Escalar al mÃ¡ximo permitido
        width_scale = max_width / float(img.width)
        height_scale = max_height / float(img.height)
        scale = min(width_scale, height_scale)
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)

        # Calcular tamaÃ±o del header en px
        header_height_px = 0
        default_height = getattr(ws.sheet_format, "defaultRowHeight", 15.0)
        for r in range(r1, r2 + 1):
            header_height_px += self._row_height_to_pixels(ws.row_dimensions[r].height, default_height)

        header_width_px = 0
        for c in range(c1, c2 + 1):
            col_letter = get_column_letter(c)
            header_width_px += self._column_width_to_pixels(ws.column_dimensions[col_letter].width)

        x_off_px = max(int(round((header_width_px - img.width) / 2)), 0)
        y_off_px = max(int(round((header_height_px - img.height) / 2)), 0)

        # EMU (openpyxl)
        emu_per_px = 9525
        marker = AnchorMarker(col=c1 - 1, colOff=x_off_px * emu_per_px, row=r1 - 1, rowOff=y_off_px * emu_per_px)
        ext = XDRPositiveSize2D(img.width * emu_per_px, img.height * emu_per_px)
        img.anchor = OneCellAnchor(_from=marker, ext=ext)
        ws.add_image(img)
