from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *
from backend.shared.tabular.pandas_mapper import from_pandas_table

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class MetricsExtractMixin:
    def extract_consultant_data(
        self,
        dataframe: pd.DataFrame,
        metadata: Dict[str, object],
        date_column: str = "Fecha",
        hours_column: str = "Horas",
        hour_type_column: Optional[str] = "Tipo Hora",
    ) -> ConsultorMetrics:
        """
        Extrae y calcula métricas de un consultor desde su DataFrame procesado.

        Args:
            dataframe: DataFrame con actividades del consultor (ya validado/corregido)
            metadata: Metadata extraída del archivo
            date_column: Nombre de la columna de fecha
            hours_column: Nombre de la columna de horas
            hour_type_column: Nombre de la columna de tipo de hora (HN, HE, etc.)

        Returns:
            ConsultorMetrics con todos los cálculos
        """
        dataframe = dataframe.copy()
        self._normalize_ticket_columns(dataframe)
        self._fill_empty_special_columns(dataframe)

        metadata = dict(metadata or {})
        # Extraer nombre del consultor
        nombre = str(metadata.get("employee", "Consultor Desconocido"))

        # Inferir cargo (puede mejorarse con lógica adicional o desde metadata)
        cargo = self._inferir_cargo(nombre, metadata)
        
        # Obtener manager de tarifas
        rates_manager = self._require_collaborator_rates()
        collab = rates_manager.find_collaborator(nombre)
        if collab and collab.nombre_completo:
            nombre = collab.nombre_completo
            metadata["employee"] = nombre

        # Parsear fechas y horas
        parsed_dates = pd.to_datetime(
            dataframe[date_column], errors="coerce", dayfirst=True
        )
        numeric_hours = pd.to_numeric(dataframe[hours_column], errors="coerce")

        # Filtrar solo filas válidas
        valid_mask = parsed_dates.notna() & numeric_hours.notna() & (numeric_hours > 0)
        valid_data = dataframe[valid_mask].copy()
        valid_dates = parsed_dates[valid_mask]
        valid_hours = numeric_hours[valid_mask]

        # Calcular días laborados (fechas únicas)
        dias_laborados = len(valid_dates.dt.date.unique())

        # Calcular horas normales y extras
        total_horas_normales = 0.0
        total_horas_extras = 0.0

        if hour_type_column and hour_type_column in dataframe.columns:
            hour_types = valid_data[hour_type_column].fillna("HN").astype(str)
            for hour_type, hours in zip(hour_types, valid_hours):
                if "HN" in hour_type.upper():
                    total_horas_normales += hours
                else:
                    total_horas_extras += hours
        else:
            # Si no hay columna de tipo de hora, asumir todas como normales
            total_horas_normales = float(valid_hours.sum())

        # Obtener tarifas personalizadas del colaborador
        # Si el colaborador está en collaborator_rates.json, usa sus datos completos
        # Si no está, usa defaults por seniority (ya inferido por _inferir_cargo)
        rate_info = rates_manager.get_rate_for_collaborator(nombre, seniority_fallback=cargo)
        if rate_info.get("cargo"):
            metadata["area"] = "CD"
        else:
            metadata["area"] = "TI"
        
        # Usar valores específicos del colaborador si están disponibles
        valor_tarifa = rate_info["salario_mensual"]
        valor_hora_extra = rate_info["valor_hora_extra"]
        # Días laborables: prioridad → rates.json → metadata → cálculo automático
        dias_laborables = rate_info.get("dias_laborables_mes")

        if not dias_laborables:
            dias_laborables = self._calcular_dias_laborables(
                metadata.get("year"),
                metadata.get("month_name"),
            )
            logger.info(
                "Dias laborables calculados automaticamente para %s: %d",
                nombre,
                dias_laborables,
            )

        # Valor por día SIEMPRE se recalcula de forma segura
        valor_dia = valor_tarifa / dias_laborables

        # Log si se encontró configuración personalizada
        if rate_info["found_by_name"]:
            logger.info(
                f"Usando tarifas personalizadas para {rate_info['nombre_completo']}: "
                f"${valor_dia:.2f}/día, ${valor_hora_extra:.2f}/hora extra"
            )
        else:
            logger.info(
                f"Usando tarifas por seniority ({cargo}) para {nombre}: "
                f"${valor_dia:.2f}/día"
            )
        total_facturar = valor_dia * dias_laborados
        total_horas_extras_facturar = total_horas_extras * valor_hora_extra

        return ConsultorMetrics(
            nombre=nombre,
            cargo=cargo,
            dias_laborados=dias_laborados,
            total_horas_normales=total_horas_normales,
            total_horas_extras=total_horas_extras,
            valor_tarifa=valor_tarifa,
            valor_dia=valor_dia,
            total_facturar=total_facturar,
            valor_hora_extra=valor_hora_extra,
            total_horas_extras_facturar=total_horas_extras_facturar,
            dataframe=from_pandas_table(dataframe),
            metadata=metadata,
        )

    @staticmethod
    def _normalize_ticket_columns(df: pd.DataFrame) -> None:
        ticket_columns = [col for col in df.columns if "ticket" in str(col).lower()]
        if not ticket_columns:
            return

        def _normalize(value: object) -> object:
            if pd.isna(value):
                return ""
            if isinstance(value, (int, float)):
                if isinstance(value, float) and value.is_integer():
                    return str(int(value))
                return str(value)
            text = str(value).strip()
            if not text:
                return ""
            normalized = re.sub(r"\s+", "", text)
            return normalized

        for col in ticket_columns:
            df[col] = df[col].apply(_normalize)

    @staticmethod
    def _excel_safe_value(value: object) -> object:
        """Evita que openpyxl falle con valores <NA> o NaN."""
        if value is None:
            return None
        try:
            if pd.isna(value):
                return None
        except Exception:
            pass
        return value

    @staticmethod
    def _get_used_range_limits(ws: Worksheet) -> Tuple[int, int]:
        max_row = 0
        max_col = 0
        for cell in ws._cells.values():
            value = cell.value
            if value is None:
                continue
            if isinstance(value, str) and not value.strip():
                continue
            if cell.row > max_row:
                max_row = cell.row
            if cell.column > max_col:
                max_col = cell.column
        for merged in ws.merged_cells.ranges:
            if merged.max_row > max_row:
                max_row = merged.max_row
            if merged.max_col > max_col:
                max_col = merged.max_col
        if max_row == 0:
            max_row = ws.max_row or 1
        if max_col == 0:
            max_col = ws.max_column or 1
        return max_row, max_col

    @staticmethod
    def _row_is_empty(ws: Worksheet, row: int, max_col: int) -> bool:
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            value = cell.value
            if value is None:
                continue
            if isinstance(value, str) and not value.strip():
                continue
            return False
        return True

    @staticmethod
    def _row_has_merge(merged_ranges: List[object], row: int) -> bool:
        for merged in merged_ranges:
            if merged.min_row <= row <= merged.max_row:
                return True
        return False

    @staticmethod
    def _clear_row_format(ws: Worksheet, row: int, max_col: int) -> None:
        ws.row_dimensions[row].height = None
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                continue
            cell.border = Border()
            cell.fill = PatternFill()
            cell.alignment = Alignment()
            cell.number_format = "General"
