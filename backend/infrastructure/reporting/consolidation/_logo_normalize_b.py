from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class LogoNormalizeBMixin:
    def normalize_header_logos(
        self,
        ws: Worksheet,
        bit_logo_path: Optional[Path],
        nova_logo_path: Optional[Path],
        *,
        title_text: str = "Informe de Actividades",
        bit_max_size: Tuple[int, int] = (150, 52),
        nova_max_size: Tuple[int, int] = (140, 48),
        margin_left_px: int = 18,
        margin_right_px: int = 18,
        safe_gap_px: int = 12,
    ) -> None:
        """Wrapper legacy: usa tamaños fijos (px) y normaliza con EMU."""
        if not bit_logo_path or not nova_logo_path:
            return
        bit_target = self._compute_target_size(bit_logo_path, *bit_max_size)
        nova_target = self._compute_target_size(nova_logo_path, *nova_max_size)
        if not bit_target or not nova_target:
            return
        emu_per_px = 9525
        ref_bit = {"cx": int(bit_target[0]) * emu_per_px, "cy": int(bit_target[1]) * emu_per_px}
        ref_nova = {"cx": int(nova_target[0]) * emu_per_px, "cy": int(nova_target[1]) * emu_per_px}
        self.normalize_sheet_header_logos(
            ws,
            ref_bit,
            ref_nova,
            bit_logo_path,
            nova_logo_path,
            title_texts=[title_text, "Informe Consolidado de Actividades"],
        )

    def _position_logo_in_block(
        self,
        ws: Worksheet,
        img_path: Path,
        header_rows: Tuple[int, int],
        block_cols: Tuple[int, int],
        *,
        max_width: int,
        max_height: int,
    ) -> None:
        if not img_path.exists():
            return
        img = Image(str(img_path))
        if not img.width or not img.height:
            return

        r1, r2 = header_rows
        c1, c2 = block_cols

        header_height_px = 0
        default_height = getattr(ws.sheet_format, "defaultRowHeight", 15.0)
        for r in range(r1, r2 + 1):
            header_height_px += self._row_height_to_pixels(ws.row_dimensions[r].height, default_height)

        block_width_px = 0
        for c in range(c1, c2 + 1):
            col_letter = get_column_letter(c)
            block_width_px += self._column_width_to_pixels(ws.column_dimensions[col_letter].width)

        # Escalado anti-recorte
        scale = min(
            block_width_px / float(img.width) if block_width_px else 1.0,
            header_height_px / float(img.height) if header_height_px else 1.0,
            max_width / float(img.width),
            max_height / float(img.height),
        )
        scale = min(scale, 1.0)
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)

        x_off_px = max(int(round((block_width_px - img.width) / 2)), 0)
        y_off_px = max(int(round((header_height_px - img.height) / 2)), 0)

        emu_per_px = 9525
        marker = AnchorMarker(col=c1 - 1, colOff=x_off_px * emu_per_px, row=r1 - 1, rowOff=y_off_px * emu_per_px)
        ext = XDRPositiveSize2D(img.width * emu_per_px, img.height * emu_per_px)
        img.anchor = OneCellAnchor(_from=marker, ext=ext)
        ws.add_image(img)

    def get_header_layout(
        self,
        ws: Worksheet,
        title_text: str = "Informe de Actividades",
    ) -> Dict[str, object]:
        """Detecta layout del header y calcula cajas para logos."""
        title_merge = self._find_title_merge(ws, title_text)
        if title_merge:
            header_rows = (1, title_merge[1])
        else:
            header_rows = (1, min(3, ws.max_row or 3))

        last_col = self._get_last_header_col(ws, header_rows, title_merge)

        def _box_width_px(c1: int, c2: int) -> int:
            width = 0
            for c in range(c1, c2 + 1):
                col_letter = get_column_letter(c)
                width += self._column_width_to_pixels(ws.column_dimensions[col_letter].width)
            return width

        header_height_px = 0
        default_height = getattr(ws.sheet_format, "defaultRowHeight", 15.0)
        for r in range(header_rows[0], header_rows[1] + 1):
            header_height_px += self._row_height_to_pixels(ws.row_dimensions[r].height, default_height)

        if title_merge:
            left_box = (1, max(title_merge[2] - 1, 1))
            right_box = (min(title_merge[3] + 1, last_col), last_col)
        else:
            left_box = (1, max(1, min(3, last_col)))
            right_box = (max(1, last_col - 1), last_col)

        min_box_px = 60
        left_w = _box_width_px(left_box[0], left_box[1]) if left_box[1] >= left_box[0] else 0
        right_w = _box_width_px(right_box[0], right_box[1]) if right_box[1] >= right_box[0] else 0

        if left_w < min_box_px or right_w < min_box_px:
            # Fallback por porcentaje
            left_end = max(1, int(round(last_col * 0.2)))
            right_start = max(left_end + 1, int(round(last_col * 0.8)))
            left_box = (1, left_end)
            right_box = (right_start, last_col)

        return {
            "header_rows": header_rows,
            "title_merge_range": title_merge,
            "lastColUsedInHeader": last_col,
            "leftLogoBox": left_box,
            "rightLogoBox": right_box,
            "headerHeightPx": header_height_px,
        }

    def place_logo_in_box(
        self,
        ws: Worksheet,
        img_path: Path,
        box: Tuple[int, int],
        header_rows: Tuple[int, int],
        max_scale: float = 1.0,
    ) -> None:
        """Posiciona logo dentro de un box con escalado y offsets reales."""
        if not img_path.exists():
            return

        if PILImage:
            with PILImage.open(img_path) as pil_img:
                img_w = pil_img.width
                img_h = pil_img.height
        else:
            img = Image(str(img_path))
            img_w = img.width
            img_h = img.height

        if not img_w or not img_h:
            return

        r1, r2 = header_rows
        c1, c2 = box

        header_height_px = 0
        default_height = getattr(ws.sheet_format, "defaultRowHeight", 15.0)
        for r in range(r1, r2 + 1):
            header_height_px += self._row_height_to_pixels(ws.row_dimensions[r].height, default_height)

        box_width_px = 0
        for c in range(c1, c2 + 1):
            col_letter = get_column_letter(c)
            box_width_px += self._column_width_to_pixels(ws.column_dimensions[col_letter].width)

        if box_width_px == 0 or header_height_px == 0:
            return

        scale = min(box_width_px / float(img_w), header_height_px / float(img_h)) * 0.92
        if max_scale is not None:
            scale = min(scale, max_scale)

        new_w = int(img_w * scale)
        new_h = int(img_h * scale)

        x_off_px = max(int(round((box_width_px - new_w) / 2)), 0)
        y_off_px = max(int(round((header_height_px - new_h) / 2)), 0)

        img = Image(str(img_path))
        img.width = new_w
        img.height = new_h

        emu_per_px = 9525
        marker = AnchorMarker(col=c1 - 1, colOff=x_off_px * emu_per_px, row=r1 - 1, rowOff=y_off_px * emu_per_px)
        ext = XDRPositiveSize2D(new_w * emu_per_px, new_h * emu_per_px)
        img.anchor = OneCellAnchor(_from=marker, ext=ext)
        ws.add_image(img)

    def normalize_header(
        self,
        ws: Worksheet,
        bit_logo_path: Optional[Path],
        nova_logo_path: Optional[Path],
        title_text: str = "Informe de Actividades",
    ) -> None:
        """Normaliza logos BIT/NOVA con layout detectado del header."""
        layout = self.get_header_layout(ws, title_text=title_text)
        header_rows = layout["header_rows"]
        left_box = layout["leftLogoBox"]
        right_box = layout["rightLogoBox"]
        last_col = layout["lastColUsedInHeader"]

        # Eliminar logos existentes en header (izq/der)
        kept_images = []
        for img in getattr(ws, "_images", []):
            anchor = img.anchor
            img_col = None
            img_row = None
            if isinstance(anchor, str):
                try:
                    col_letter, row = re.match(r"([A-Z]+)(\d+)", anchor).groups()
                    img_col = column_index_from_string(col_letter)
                    img_row = int(row)
                except Exception:
                    img_col = None
            elif hasattr(anchor, "_from"):
                img_col = anchor._from.col + 1
                img_row = anchor._from.row + 1

            if img_col is not None and img_row is not None:
                if header_rows[0] <= img_row <= header_rows[1]:
                    if img_col <= int(last_col * 0.3) or img_col >= int(last_col * 0.7):
                        continue
            kept_images.append(img)
        ws._images = kept_images

        if bit_logo_path and bit_logo_path.exists():
            self.place_logo_in_box(ws, bit_logo_path, left_box, header_rows, max_scale=1.0)
        if nova_logo_path and nova_logo_path.exists():
            self.place_logo_in_box(ws, nova_logo_path, right_box, header_rows, max_scale=1.0)
