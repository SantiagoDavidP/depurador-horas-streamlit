from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from backend.collaborator_rates import get_collaborator_rates_manager

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class LogoHeaderMixin:
    def _find_title_merge(
        self,
        ws: Worksheet,
        title_text: str,
    ) -> Optional[Tuple[int, int, int, int]]:
        """Encuentra el rango merged que contiene el tÃ­tulo."""
        target = title_text.lower()
        title_cell = None
        for row in ws.iter_rows(min_row=1, max_row=8):
            for cell in row:
                if cell.value and target in str(cell.value).lower():
                    title_cell = cell
                    break
            if title_cell:
                break
        if not title_cell:
            return None
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= title_cell.row <= mr.max_row and mr.min_col <= title_cell.column <= mr.max_col:
                return (mr.min_row, mr.max_row, mr.min_col, mr.max_col)
        return (title_cell.row, title_cell.row, title_cell.column, title_cell.column)

    def _find_title_merge_any(
        self,
        ws: Worksheet,
        title_texts: List[str],
    ) -> Optional[Tuple[int, int, int, int]]:
        """Encuentra el rango merged que contiene cualquiera de los tÃƒÂ­tulos."""
        targets = [t.lower() for t in title_texts]
        title_cell = None
        for row in ws.iter_rows(min_row=1, max_row=8):
            for cell in row:
                if cell.value:
                    cell_text = str(cell.value).lower()
                    if any(t in cell_text for t in targets):
                        title_cell = cell
                        break
            if title_cell:
                break
        if not title_cell:
            return None
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= title_cell.row <= mr.max_row and mr.min_col <= title_cell.column <= mr.max_col:
                return (mr.min_row, mr.max_row, mr.min_col, mr.max_col)
        return (title_cell.row, title_cell.row, title_cell.column, title_cell.column)

    def _get_last_header_col(
        self,
        ws: Worksheet,
        header_rows: Tuple[int, int],
        title_merge: Optional[Tuple[int, int, int, int]],
    ) -> int:
        """Detecta la Ãºltima columna usada en el header."""
        r1, r2 = header_rows
        last_col = title_merge[3] if title_merge else ws.max_column or 1
        max_col_scan = ws.max_column or last_col
        for c in range(1, max_col_scan + 1):
            for r in range(r1, r2 + 1):
                cell = ws.cell(row=r, column=c)
                if cell.value is not None:
                    last_col = max(last_col, c)
                if cell.fill and cell.fill.fill_type:
                    last_col = max(last_col, c)
        return last_col

    def get_header_rows(
        self,
        ws: Worksheet,
        title_texts: Optional[List[str]] = None,
    ) -> Tuple[Tuple[int, int], int, Optional[Tuple[int, int, int, int]]]:
        """Devuelve filas del header, alto en EMU y rango del título."""
        if not title_texts:
            title_texts = ["Informe de Actividades", "Informe Consolidado de Actividades"]
        title_merge = self._find_title_merge_any(ws, title_texts)
        if title_merge:
            title_min_row, title_max_row, title_min_col, title_max_col = title_merge
            header_start = 1
            header_end = title_max_row
            max_scan = min(title_max_row + 6, ws.max_row)
            for r in range(title_max_row + 1, max_scan + 1):
                has_fill = False
                for c in range(title_min_col, title_max_col + 1):
                    cell = ws.cell(row=r, column=c)
                    if cell.fill and cell.fill.fill_type:
                        has_fill = True
                        break
                if has_fill:
                    header_end = r
                else:
                    break
        else:
            header_start = 1
            header_end = min(3, ws.max_row or 3)

        header_rows = (header_start, header_end)
        header_height_emu = 0
        default_height = getattr(ws.sheet_format, "defaultRowHeight", 15.0)
        for r in range(header_start, header_end + 1):
            header_height_emu += self._row_height_to_emu(ws.row_dimensions[r].height, default_height)
        return header_rows, header_height_emu, title_merge

    def get_header_width(
        self,
        ws: Worksheet,
        header_rows: Tuple[int, int],
        title_merge: Optional[Tuple[int, int, int, int]],
    ) -> Tuple[int, int]:
        """Devuelve ancho del header en EMU y última columna usada."""
        last_col = self._get_last_header_col(ws, header_rows, title_merge)
        header_width_emu = 0
        for c in range(1, last_col + 1):
            col_letter = get_column_letter(c)
            header_width_emu += self._column_width_to_emu(ws.column_dimensions[col_letter].width)
        return header_width_emu, last_col

    def _get_title_x_range_emu(
        self,
        ws: Worksheet,
        title_merge: Optional[Tuple[int, int, int, int]],
        header_width_emu: int,
    ) -> Tuple[int, int]:
        if not title_merge or header_width_emu <= 0:
            return (int(round(header_width_emu * 0.35)), int(round(header_width_emu * 0.65)))
        _, _, title_min_col, title_max_col = title_merge
        start_x = 0
        for c in range(1, title_min_col):
            col_letter = get_column_letter(c)
            start_x += self._column_width_to_emu(ws.column_dimensions[col_letter].width)
        end_x = start_x
        for c in range(title_min_col, title_max_col + 1):
            col_letter = get_column_letter(c)
            end_x += self._column_width_to_emu(ws.column_dimensions[col_letter].width)
        if start_x <= header_width_emu * 0.1 and end_x >= header_width_emu * 0.9:
            return (int(round(header_width_emu * 0.35)), int(round(header_width_emu * 0.65)))
        return (start_x, end_x)

    def xy_to_anchor(
        self,
        ws: Worksheet,
        x_emu: int,
        y_emu: int,
        start_row: int = 1,
    ) -> Tuple[int, int, int, int]:
        """Convierte x/y en EMU a (col,row,colOff,rowOff) 0-based."""
        x_emu = max(int(x_emu), 0)
        y_emu = max(int(y_emu), 0)

        col_idx = 1
        x_acc = 0
        while True:
            col_letter = get_column_letter(col_idx)
            col_w = self._column_width_to_emu(ws.column_dimensions[col_letter].width)
            if x_acc + col_w > x_emu or col_idx >= (ws.max_column or 1):
                break
            x_acc += col_w
            col_idx += 1
        col_off = max(x_emu - x_acc, 0)

        row_idx = start_row
        y_acc = 0
        default_height = getattr(ws.sheet_format, "defaultRowHeight", 15.0)
        max_row = ws.max_row or start_row
        while True:
            row_h = self._row_height_to_emu(ws.row_dimensions[row_idx].height, default_height)
            if y_acc + row_h > y_emu or row_idx >= max_row:
                break
            y_acc += row_h
            row_idx += 1
        row_off = max(y_emu - y_acc, 0)

        return col_idx - 1, row_idx - 1, int(col_off), int(row_off)

    def _image_anchor_xy_emu(
        self,
        ws: Worksheet,
        img: Image,
    ) -> Optional[Tuple[int, int, int]]:
        anchor = img.anchor
        emu_per_px = 9525
        default_height = getattr(ws.sheet_format, "defaultRowHeight", 15.0)

        def _col_start(col_idx: int) -> int:
            total = 0
            for c in range(1, col_idx):
                col_letter = get_column_letter(c)
                total += self._column_width_to_emu(ws.column_dimensions[col_letter].width)
            return total

        def _row_start(row_idx: int) -> int:
            total = 0
            for r in range(1, row_idx):
                total += self._row_height_to_emu(ws.row_dimensions[r].height, default_height)
            return total

        if isinstance(anchor, str):
            try:
                col_letter, row = re.match(r"([A-Z]+)(\d+)", anchor).groups()
                col_idx = column_index_from_string(col_letter)
                row_idx = int(row)
                x_emu = _col_start(col_idx)
                y_emu = _row_start(row_idx)
                return (x_emu, y_emu, row_idx)
            except Exception:
                return None
        if hasattr(anchor, "_from"):
            try:
                col_idx = anchor._from.col + 1
                row_idx = anchor._from.row + 1
                col_off = int(getattr(anchor._from, "colOff", 0))
                row_off = int(getattr(anchor._from, "rowOff", 0))
                x_emu = _col_start(col_idx) + col_off
                y_emu = _row_start(row_idx) + row_off
                return (x_emu, y_emu, row_idx)
            except Exception:
                return None
        if hasattr(anchor, "col") and hasattr(anchor, "row"):
            try:
                col_idx = int(anchor.col) + 1
                row_idx = int(anchor.row) + 1
                x_emu = _col_start(col_idx)
                y_emu = _row_start(row_idx)
                return (x_emu, y_emu, row_idx)
            except Exception:
                return None
        # Fallback con imagen en px si no hay anchor
        if img.width and img.height:
            return (0, 0, 1)
        return None
