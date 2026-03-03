from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from backend.domain.rates.collaborator_rates import get_collaborator_rates_manager

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class FooterBuilderMixin:
    def buildFooter(
        self,
        ws: Worksheet,
        startRow: int,
        resourceName: str,
        approverName: str,
        approverTitle: str,
        sharepointPath: str,
        firstDetailRow: int,
        lastDetailRow: int,
    ) -> None:
        """Construye el pie con formato Alfredo Aguirre."""
        header_row = max(1, firstDetailRow - 1)
        tipo_col = None
        horas_col = None
        for col_idx in range(1, ws.max_column + 1):
            header_value = ws.cell(row=header_row, column=col_idx).value
            if not header_value:
                continue
            header_text = str(header_value).strip().lower()
            if "tipo hora" in header_text:
                tipo_col = col_idx
            elif "horas" in header_text or header_text == "hora":
                if "tipo" not in header_text:
                    horas_col = col_idx

        tipo_letter = get_column_letter(tipo_col) if tipo_col else None
        horas_letter = get_column_letter(horas_col) if horas_col else None

        def _find_last_data_row() -> int:
            if not tipo_col and not horas_col:
                return lastDetailRow
            for row in range(lastDetailRow, firstDetailRow - 1, -1):
                tipo_val = ws.cell(row=row, column=tipo_col).value if tipo_col else None
                horas_val = ws.cell(row=row, column=horas_col).value if horas_col else None
                if (tipo_val is not None and str(tipo_val).strip() != "") or (
                    horas_val is not None and str(horas_val).strip() != ""
                ):
                    return row
            return lastDetailRow

        def _find_extra_start_row(last_row: int) -> int:
            if not tipo_col:
                return firstDetailRow
            for row in range(firstDetailRow, last_row + 1):
                value = ws.cell(row=row, column=tipo_col).value
                if value is None:
                    continue
                if str(value).strip().upper() == "HS":
                    return row
            return firstDetailRow

        def _sumif(tipo: str) -> str:
            if not tipo_letter or not horas_letter:
                return ""
            last_row = _find_last_data_row()
            tipo_range = f"${tipo_letter}${firstDetailRow}:${tipo_letter}${last_row}"
            horas_range = f"${horas_letter}${firstDetailRow}:${horas_letter}${last_row}"
            return f'=SUMIF({tipo_range},"{tipo}",{horas_range})'

        def _sumif_hs() -> str:
            if not tipo_letter or not horas_letter:
                return ""
            last_row = _find_last_data_row()
            extra_start_row = _find_extra_start_row(last_row)
            tipo_range = f"${tipo_letter}${firstDetailRow}:${tipo_letter}${last_row}"
            horas_range = f"${horas_letter}${extra_start_row}:${horas_letter}${last_row}"
            return f'=SUMIF({tipo_range},"HS",{horas_range})'

        bold_font = Font(name="Calibri", size=10, bold=True, color=COLOR_BLACK)
        normal_font = Font(name="Calibri", size=10, bold=False, color=COLOR_BLACK)
        center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
        gray_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

        r_total_1 = startRow
        r_total_2 = startRow + 1
        r_total_3 = startRow + 2
        r_legend = startRow + 4
        r_repo = startRow + 6
        r_sig_header = startRow + 8
        r_sig_space = startRow + 9
        r_sig_name = startRow + 10
        r_sig_title = startRow + 11

        ws.merge_cells(f"H{r_total_1}:K{r_total_1}")
        ws.merge_cells(f"H{r_total_2}:K{r_total_2}")
        ws.merge_cells(f"H{r_total_3}:K{r_total_3}")

        ws[f"H{r_total_1}"].value = "Total Horas Normales"
        ws[f"H{r_total_2}"].value = "Total Horas Extras entre semana"
        ws[f"H{r_total_3}"].value = "Total Horas Extras fin de semana o feriado"

        ws[f"L{r_total_1}"].value = _sumif("HN")
        ws[f"L{r_total_2}"].value = _sumif_hs()
        ws[f"L{r_total_3}"].value = _sumif("HF")

        for r in (r_total_1, r_total_2, r_total_3):
            ws[f"H{r}"].font = bold_font
            ws[f"H{r}"].alignment = left_wrap
            ws[f"L{r}"].font = bold_font
            ws[f"L{r}"].alignment = center_wrap

        ws.merge_cells(f"A{r_legend}:L{r_legend}")
        ws[f"A{r_legend}"].value = (
            "Tipo Hora - HN: hora normal / HS: hora extra entre semana / "
            "HF: hora extra fin de semana o feriado"
        )
        ws[f"A{r_legend}"].font = bold_font
        ws[f"A{r_legend}"].alignment = center_wrap
        ws.row_dimensions[r_legend].height = 13

        ws.merge_cells(f"A{r_repo}:L{r_repo}")
        ws[f"A{r_repo}"].value = (
            "Repositorio Base Sharepoint de DocumentaciónEntregada:\n"
            f"{sharepointPath}"
        )
        ws[f"A{r_repo}"].font = bold_font
        ws[f"A{r_repo}"].alignment = center_wrap
        ws.row_dimensions[r_repo].height = 26

        ws.merge_cells(f"A{r_sig_header}:F{r_sig_header}")
        ws.merge_cells(f"G{r_sig_header}:L{r_sig_header}")
        ws[f"A{r_sig_header}"].value = "Elaborado por:"
        ws[f"G{r_sig_header}"].value = "Aprobado por:"
        for cell in (ws[f"A{r_sig_header}"], ws[f"G{r_sig_header}"]):
            cell.font = bold_font
            cell.alignment = center_wrap
            cell.fill = gray_fill

        ws.merge_cells(f"A{r_sig_space}:F{r_sig_space}")
        ws.merge_cells(f"G{r_sig_space}:L{r_sig_space}")
        ws.row_dimensions[r_sig_space].height = 49.5

        ws.merge_cells(f"A{r_sig_name}:F{r_sig_name}")
        ws.merge_cells(f"G{r_sig_name}:L{r_sig_name}")
        ws[f"A{r_sig_name}"].value = resourceName
        ws[f"G{r_sig_name}"].value = approverName
        ws[f"A{r_sig_name}"].font = normal_font
        ws[f"G{r_sig_name}"].font = normal_font
        ws[f"A{r_sig_name}"].alignment = center_wrap
        ws[f"G{r_sig_name}"].alignment = center_wrap

        ws.merge_cells(f"A{r_sig_title}:F{r_sig_title}")
        ws.merge_cells(f"G{r_sig_title}:L{r_sig_title}")
        ws[f"A{r_sig_title}"].value = "Cargo Recurso"
        ws[f"G{r_sig_title}"].value = approverTitle
        ws[f"A{r_sig_title}"].font = normal_font
        ws[f"G{r_sig_title}"].font = normal_font
        ws[f"A{r_sig_title}"].alignment = center_wrap
        ws[f"G{r_sig_title}"].alignment = center_wrap

        for r in range(r_total_1, r_sig_title + 1):
            for c in range(1, 13):
                ws.cell(row=r, column=c).border = THIN_BORDER
