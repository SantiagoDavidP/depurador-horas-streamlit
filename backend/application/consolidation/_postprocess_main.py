from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from backend.domain.rates.collaborator_rates import get_collaborator_rates_manager

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class PostprocessMainMixin:
    def postprocess_hide_empty_table_rows_preserve_footer(
        self, wb: Workbook, sheet_df_map: Dict[str, pd.DataFrame]
    ) -> None:
        def norm(text: object) -> str:
            return str(text or "").replace("\u200b", "").replace("\xa0", " ").strip().lower()

        def strip_invisible(text: str) -> str:
            return text.replace("\u200b", "").replace("\xa0", " ").strip()

        def detect_table_start_row(ws: Worksheet) -> Optional[int]:
            best_row = None
            best_score = -1
            for row in range(1, min(35, ws.max_row) + 1):
                values = [norm(cell.value) for cell in ws[row] if norm(cell.value)]
                if not values:
                    continue
                has_fecha = any("fecha" in v for v in values)
                has_act = any("actividad" in v or "tarea" in v for v in values)
                has_horas_or_proy = any("hora" in v or "proyecto" in v for v in values)
                score = sum(
                    1
                    for t in ("fecha", "proyecto", "horas", "actividad", "tareas")
                    if any(t in v for v in values)
                )
                if has_fecha and has_act and has_horas_or_proy:
                    return row
                if score > best_score and score >= 2:
                    best_score = score
                    best_row = row
            return best_row

        def detect_footer_row(ws: Worksheet, start_row: int) -> Optional[int]:
            for row in range(start_row + 1, ws.max_row + 1):
                val = norm(ws.cell(row=row, column=1).value)
                if val == "elaborado por:":
                    return row
            return None

        def detect_lock_row(ws: Worksheet, start_row: int, footer_row: int, scan_cols: int) -> int:
            keywords = [
                "total horas normales",
                "total horas extras",
                "tipo hora - hn",
                "repositorio base sharepoint",
            ]
            for row in range(start_row + 1, footer_row):
                for col in range(1, scan_cols + 1):
                    text = norm(ws.cell(row=row, column=col).value)
                    if not text:
                        continue
                    for kw in keywords:
                        if kw in text:
                            return row
            return footer_row

        def formula_returns_empty(formula: str) -> bool:
            if '""' not in formula:
                return False
            cleaned = formula.replace('""', "")
            return not re.search(r'"[^"]+"', cleaned)

        def is_effectively_empty(cell) -> bool:
            value = cell.value
            if value is None:
                return True
            if isinstance(value, str):
                stripped = strip_invisible(value)
                if stripped == "":
                    return True
                if cell.data_type == "f" or stripped.startswith("="):
                    cleaned = re.sub(r"\s+", "", stripped)
                    if cleaned == '=""' or formula_returns_empty(cleaned):
                        return True
                return False
            if cell.data_type == "f":
                if value is None:
                    return True
            return False

        def is_total_formula(cell) -> bool:
            value = cell.value
            if value is None:
                return False
            text = value if isinstance(value, str) else ""
            if not text and cell.data_type == "f":
                text = str(value)
            if not text:
                return False
            upper = text.upper().replace(" ", "")
            return "SUM(" in upper or "SUMIF(" in upper or "SUBTOTAL(" in upper

        def is_day_only(value: object) -> bool:
            if value is None:
                return False
            if isinstance(value, (date, datetime)):
                return False
            if isinstance(value, (int, float)):
                try:
                    return 1 <= int(value) <= 31
                except Exception:
                    return False
            if isinstance(value, str):
                raw = strip_invisible(value)
                if raw.isdigit():
                    try:
                        num = int(raw)
                        return 1 <= num <= 31
                    except Exception:
                        return False
                if "/" in raw or "-" in raw:
                    return False
                try:
                    num = int(float(raw))
                    return 1 <= num <= 31
                except Exception:
                    return False
            return False

        for ws in wb.worksheets:
            if "resumen" in ws.title.lower():
                continue
            try:
                table_start_row = detect_table_start_row(ws)
                if not table_start_row:
                    logger.info("hide_empty: %s sin header detectado, omitido", ws.title)
                    continue

                footer_row = detect_footer_row(ws, table_start_row)
                if not footer_row:
                    logger.info("hide_empty: %s sin footer detectado, omitido", ws.title)
                    continue

                # Determinar última columna con header no vacío
                table_max_col = 0
                for cell in ws[table_start_row]:
                    if norm(cell.value):
                        table_max_col = max(table_max_col, cell.column)
                if table_max_col == 0:
                    logger.info("hide_empty: %s header vacío, omitido", ws.title)
                    continue

                # Determinar columnas clave
                header_values = [
                    norm(ws.cell(row=table_start_row, column=c).value)
                    for c in range(1, table_max_col + 1)
                ]
                col_fecha = next((i for i, v in enumerate(header_values, start=1) if "fecha" in v), None)
                col_actividad = next(
                    (i for i, v in enumerate(header_values, start=1) if "actividad" in v or "tarea" in v or "tareas" in v),
                    None,
                )
                col_horas = next((i for i, v in enumerate(header_values, start=1) if "hora" in v), None)

                if not col_fecha or not col_actividad or not col_horas:
                    logger.info(
                        "hide_empty: %s sin col_fecha/col_actividad/col_horas, omitido",
                        ws.title,
                    )
                    continue

                scan_cols = max(table_max_col, 12)
                lock_row = detect_lock_row(ws, table_start_row, footer_row, scan_cols)
                rows_in_df = len(sheet_df_map.get(ws.title, []))

                # Snapshot footer para validacion
                footer_max_row = min(lock_row + 120, ws.max_row)
                footer_max_col = max(table_max_col, 12)
                footer_cells: Dict[Tuple[int, int], Tuple[object, str, object, object, object, object]] = {}
                footer_rows_meta: Dict[int, Tuple[Optional[float], Optional[bool]]] = {}
                for r in range(lock_row, footer_max_row + 1):
                    rd = ws.row_dimensions.get(r)
                    footer_rows_meta[r] = (
                        rd.height if rd else None,
                        rd.hidden if rd else None,
                    )
                    for c in range(1, footer_max_col + 1):
                        cell = ws.cell(row=r, column=c)
                        footer_cells[(r, c)] = (
                            cell.value,
                            cell.number_format,
                            copy(cell.font),
                            copy(cell.fill),
                            copy(cell.border),
                            copy(cell.alignment),
                        )
                footer_merges = {
                    str(mr)
                    for mr in ws.merged_cells.ranges
                    if not (mr.max_row < lock_row or mr.min_row > footer_max_row)
                }

                hidden_rows: List[int] = []
                protected_rows: List[int] = []
                for r in range(table_start_row + 1, lock_row):
                    horas_cell = ws.cell(row=r, column=col_horas)
                    if is_total_formula(horas_cell):
                        protected_rows.append(r)
                        continue

                    fecha_empty = is_effectively_empty(ws.cell(row=r, column=col_fecha))
                    act_empty = is_effectively_empty(ws.cell(row=r, column=col_actividad))
                    horas_empty = is_effectively_empty(horas_cell)

                    if fecha_empty and act_empty and horas_empty:
                        ws.row_dimensions[r].hidden = True
                        ws.row_dimensions[r].height = 1
                        for c in range(1, table_max_col + 1):
                            cell = ws.cell(row=r, column=c)
                            if is_effectively_empty(cell):
                                cell.border = Border()
                                cell.fill = PatternFill()
                                cell.font = Font()
                                cell.alignment = Alignment()
                                cell.number_format = "General"
                        hidden_rows.append(r)

                # Validar footer intacto
                footer_ok = True
                for r in range(lock_row, footer_max_row + 1):
                    rd = ws.row_dimensions.get(r)
                    current_meta = (
                        rd.height if rd else None,
                        rd.hidden if rd else None,
                    )
                    if footer_rows_meta.get(r) != current_meta:
                        footer_ok = False
                        break
                    for c in range(1, footer_max_col + 1):
                        cell = ws.cell(row=r, column=c)
                        snapshot = footer_cells.get((r, c))
                        current = (
                            cell.value,
                            cell.number_format,
                            cell.font,
                            cell.fill,
                            cell.border,
                            cell.alignment,
                        )
                        if snapshot is None or current != snapshot:
                            footer_ok = False
                            break
                    if not footer_ok:
                        break
                current_merges = {
                    str(mr)
                    for mr in ws.merged_cells.ranges
                    if not (mr.max_row < lock_row or mr.min_row > footer_max_row)
                }
                if current_merges != footer_merges:
                    footer_ok = False

                logger.info(
                    "hide_empty: %s table_start=%s lock_row=%s footer_row=%s max_col=%s col_fecha=%s col_actividad=%s col_horas=%s rows_in_df=%s protected=%s hidden=%s protected_validation=%s",
                    ws.title,
                    table_start_row,
                    lock_row,
                    footer_row,
                    table_max_col,
                    col_fecha,
                    col_actividad,
                    col_horas,
                    rows_in_df,
                    protected_rows if len(protected_rows) <= 50 else f"{len(protected_rows)} rows",
                    hidden_rows if len(hidden_rows) <= 50 else f"{len(hidden_rows)} rows",
                    "OK" if footer_ok else "FAIL",
                )

                if not footer_ok:
                    raise RuntimeError(f"Footer modificado en hoja {ws.title}")
            except Exception as exc:
                logger.error("hide_empty: error en hoja %s: %s", ws.title, exc)
                raise
