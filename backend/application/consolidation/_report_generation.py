from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from backend.domain.rates.collaborator_rates import get_collaborator_rates_manager

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class ReportGenerationMixin:
    def generate_consolidated_report(
        self,
        consultores_data: List[Tuple[pd.DataFrame, Dict[str, object]]],
        output_filename: Optional[str] = None,
    ) -> ConsolidatedReport:
        """
        Genera el reporte consolidado a partir de los datos de múltiples consultores.

        Args:
            consultores_data: Lista de tuplas (dataframe, metadata) de cada consultor
            output_filename: Nombre del archivo de salida (None para auto-generar)

        Returns:
            ConsolidatedReport con el workbook y metadatos
        """
        logger.info("Iniciando generación de reporte consolidado...")

        # Extraer métricas de cada consultor
        self.consultores_metrics = []
        for df, metadata in consultores_data:
            try:
                metrics = self.extract_consultant_data(df, metadata)
                self.consultores_metrics.append(metrics)
                logger.info(
                    "Procesado: %s - %d días laborados, %.1f horas",
                    metrics.nombre,
                    metrics.dias_laborados,
                    metrics.total_horas_normales + metrics.total_horas_extras,
                )
            except Exception as exc:
                logger.error(
                    "Error procesando consultor %s: %s",
                    metadata.get("employee", "desconocido"),
                    exc,
                )
                continue

        if not self.consultores_metrics:
            raise ValueError("No se pudieron procesar consultores")

        sorted_metrics = self._get_sorted_metrics(self.consultores_metrics)

        self._resolve_cliente_from_metrics()

        # Crear workbook
        wb = Workbook()
        wb.remove(wb.active)  # Eliminar hoja por defecto

        # Determinar periodo y días laborables del primer consultor
        first_metadata = self.consultores_metrics[0].metadata
        periodo = f"{first_metadata.get('month_name', 'N/A')} {first_metadata.get('year', 'N/A')}"
        dias_laborables = self._calcular_dias_laborables(
            first_metadata.get("year"), first_metadata.get("month_name")
        )

        # 1. Crear hoja de resumen
        ws_resumen = wb.create_sheet("Resumen", 0)
        sheet_names = [
            self._sanitize_sheet_name(f"{idx}. {metrics.nombre}")
            for idx, metrics in enumerate(sorted_metrics, start=1)
        ]
        self._create_summary_sheet(
            ws_resumen, periodo, dias_laborables, sorted_metrics, sheet_names
        )

        sheet_df_map: Dict[str, pd.DataFrame] = {}

        # 2. Crear hojas individuales por consultor
        for idx, metrics in enumerate(sorted_metrics, start=1):
            sheet_name = self._sanitize_sheet_name(f"{idx}. {metrics.nombre}")
            ws_individual = wb.create_sheet(sheet_name)
            self._create_individual_sheet(ws_individual, metrics)
            sheet_df_map[sheet_name] = metrics.dataframe

        # Normalizar logos en header usando referencia del Resumen
        if self.client_logo_path and self.logo_path:
            refs = self.extract_logo_reference_from_resumen(wb)
            if refs:
                ref_bit, ref_nova = refs
                for ws in wb.worksheets:
                    self.normalize_sheet_header_logos(
                        ws,
                        ref_bit,
                        ref_nova,
                        self.logo_path,
                        self.client_logo_path,
                        title_texts=[
                            "Informe de Actividades",
                            "Informe Consolidado de Actividades",
                            "INFORME ACTIVIDADES OUTSOURCING BUSINESS IT",
                        ],
                    )
            else:
                logger.warning("No se pudo extraer referencia de logos desde Resumen.")

        # Post-procesar tablas para ocultar filas vacias (sin tocar footer)
        self.postprocess_hide_empty_table_rows_preserve_footer(wb, sheet_df_map)

        # Guardar en bytes
        from io import BytesIO

        buffer = BytesIO()
        wb.save(buffer)
        workbook_bytes = buffer.getvalue()

        # Generar nombre de archivo si no se proporciona
        if not output_filename:
            mes = first_metadata.get("month_name", "Unknown")
            year = first_metadata.get("year", "2025")
            output_filename = f"Informe_TI_Nova_BIT_{mes}_{year}_Consolidado.xlsx"

        # Calcular totales
        total_facturar = sum(
            m.total_facturar + m.total_horas_extras_facturar
            for m in self.consultores_metrics
        )
        total_horas = sum(
            m.total_horas_normales + m.total_horas_extras
            for m in self.consultores_metrics
        )

        logger.info("Reporte consolidado generado exitosamente")
        logger.info("Total consultores: %d", len(self.consultores_metrics))
        logger.info("Total a facturar: $%.2f", total_facturar)

        return ConsolidatedReport(
            workbook_bytes=workbook_bytes,
            output_filename=output_filename,
            consultores_incluidos=len(self.consultores_metrics),
            total_facturar=total_facturar,
            total_horas=total_horas,
            periodo=periodo,
            dias_laborables=dias_laborables,
        )

    def generate_single_consultant_report(
        self,
        dataframe: pd.DataFrame,
        metadata: Dict[str, object],
        output_filename: Optional[str] = None,
    ) -> bytes:
        """
        Genera un Excel individual con el mismo formato visual del consolidado.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = self._sanitize_sheet_name(str(metadata.get("employee", "Consultor")))

        metrics = self.extract_consultant_data(dataframe, metadata)
        self._create_individual_sheet(ws, metrics)

        # Limpieza visual final de tablas (no toca footer)
        self.postprocess_hide_empty_table_rows_preserve_footer(wb, {ws.title: metrics.dataframe})

        # Normalizar logos en header usando referencia del propio sheet
        if self.client_logo_path and self.logo_path:
            refs = self.extract_logo_reference_from_resumen(wb)
            if refs:
                ref_bit, ref_nova = refs
                self.normalize_sheet_header_logos(
                    ws,
                    ref_bit,
                    ref_nova,
                    self.logo_path,
                    self.client_logo_path,
                    title_texts=[
                        "Informe de Actividades",
                        "Informe Consolidado de Actividades",
                        "INFORME ACTIVIDADES OUTSOURCING BUSINESS IT",
                    ],
                )

        from io import BytesIO
        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
