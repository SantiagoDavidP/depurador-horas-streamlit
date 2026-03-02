from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from backend.collaborator_rates import get_collaborator_rates_manager

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class PostprocessHelpersMixin:
    @staticmethod
    def _detect_footer_format(footer_block: Optional[Dict[str, object]]) -> str:
        """Detecta el formato de footer basado en textos del archivo original."""
        if not footer_block:
            return "alfredo"
        cells = footer_block.get("cells") or []
        texts = []
        for cell in cells:
            value = cell.get("value")
            if value is None:
                continue
            texts.append(str(value).strip().lower())
        joined = " | ".join(texts)
        if "total horas normales" in joined or "tipo hora" in joined or "repositorio base sharepoint" in joined:
            return "alfredo"
        if "elaborado por" in joined and "aprobado por" in joined:
            return "simple"
        return "alfredo"

    @staticmethod
    def _apply_footer_block(
        ws: Worksheet,
        footer_block: Dict[str, object],
        last_data_row: int,
    ) -> None:
        """Aplica el footer original capturado desde el Excel fuente."""
        try:
            original_data_end = int(footer_block.get("data_end_row", 0))
            original_footer_start = int(footer_block.get("start_row", original_data_end + 1))
            original_footer_end = int(footer_block.get("max_row", original_footer_start))
            new_footer_start = last_data_row + 2
            offset = new_footer_start - original_footer_start

            for original_row, height in (footer_block.get("row_heights") or {}).items():
                new_row = original_row + offset
                ws.row_dimensions[new_row].height = height

            for cell_info in footer_block.get("cells", []):
                new_row = cell_info["row"] + offset
                col = cell_info["col"]
                cell = ws.cell(row=new_row, column=col, value=cell_info.get("value"))
                cell.font = cell_info.get("font")
                cell.fill = cell_info.get("fill")
                cell.border = cell_info.get("border")
                cell.alignment = cell_info.get("alignment")
                cell.number_format = cell_info.get("number_format") or cell.number_format

            for min_row, min_col, max_row, max_col in footer_block.get("merges", []):
                try:
                    ws.merge_cells(
                        start_row=min_row + offset,
                        start_column=min_col,
                        end_row=max_row + offset,
                        end_column=max_col,
                    )
                except Exception as merge_exc:
                    logger.warning("No se pudo mergear footer cells: %s", merge_exc)
        except Exception as exc:
            logger.warning("No se pudo aplicar footer original: %s", exc)

    def _resolve_cliente_from_metrics(self) -> None:
        if not self.consultores_metrics:
            return
        cd_count = 0
        ti_count = 0
        for m in self.consultores_metrics:
            area = str((m.metadata or {}).get("area", "")).upper()
            if area == "CD":
                cd_count += 1
            else:
                ti_count += 1

        # Regla por mayoría: si hay más CD que TI -> CD, si no -> TI
        if cd_count > ti_count:
            self.cliente = "NOVA - Centro Digital"
        else:
            self.cliente = "BIT Nova - TI"
