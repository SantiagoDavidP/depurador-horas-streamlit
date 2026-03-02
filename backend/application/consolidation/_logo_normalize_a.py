from __future__ import annotations

import io
import logging
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.worksheet import Worksheet

from backend.collaborator_rates import get_collaborator_rates_manager

from .models import ConsultorMetrics, ConsolidatedReport
from .styles import *

try:
    from PIL import Image as PILImage
except ImportError:
    PILImage = None  # type: ignore

logger = logging.getLogger(__name__)



class LogoNormalizeAMixin:
    def extract_logo_reference_from_resumen(
        self,
        wb: Workbook,
    ) -> Optional[Tuple[Dict[str, int], Dict[str, int]]]:
        """Extrae tamaños (cx/cy) desde la hoja Resumen para BIT y NOVA."""
        ws = wb["Resumen"] if "Resumen" in wb.sheetnames else wb.worksheets[0]
        header_rows, header_height_emu, _ = self.get_header_rows(ws)
        header_height_emu = max(header_height_emu, 1)

        header_images: List[Tuple[Image, int]] = []
        for img in getattr(ws, "_images", []):
            pos = self._image_anchor_xy_emu(ws, img)
            if not pos:
                continue
            x_emu, _, row_idx = pos
            if header_rows[0] <= row_idx <= header_rows[1]:
                header_images.append((img, x_emu))

        if len(header_images) < 2:
            return None

        header_images.sort(key=lambda item: item[1])
        left_img = header_images[0][0]
        right_img = header_images[-1][0]

        def _extract_ext(img: Image) -> Optional[Tuple[int, int]]:
            if hasattr(img.anchor, "ext") and img.anchor.ext:
                cx = int(getattr(img.anchor.ext, "cx", 0))
                cy = int(getattr(img.anchor.ext, "cy", 0))
                if cx > 0 and cy > 0:
                    return (cx, cy)
            if img.width and img.height:
                emu_per_px = 9525
                return (int(img.width) * emu_per_px, int(img.height) * emu_per_px)
            return None

        left_ext = _extract_ext(left_img)
        right_ext = _extract_ext(right_img)
        if not left_ext or not right_ext:
            return None

        return (
            {"cx": left_ext[0], "cy": left_ext[1]},
            {"cx": right_ext[0], "cy": right_ext[1]},
        )

    def normalize_sheet_header_logos(
        self,
        ws: Worksheet,
        ref_bit: Dict[str, int],
        ref_nova: Dict[str, int],
        bit_logo_path: Optional[Path],
        nova_logo_path: Optional[Path],
        *,
        title_texts: Optional[List[str]] = None,
    ) -> None:
        if not bit_logo_path or not nova_logo_path:
            return
        if not bit_logo_path.exists() or not nova_logo_path.exists():
            return

        header_rows, header_height_emu, title_merge = self.get_header_rows(ws, title_texts)
        header_width_emu, _ = self.get_header_width(ws, header_rows, title_merge)
        if header_width_emu <= 0 or header_height_emu <= 0:
            return

        title_x_start_emu, title_x_end_emu = self._get_title_x_range_emu(ws, title_merge, header_width_emu)

        # Eliminar logos existentes (izq/der) sin tocar otras imágenes
        header_images: List[Tuple[Image, int]] = []
        for img in getattr(ws, "_images", []):
            pos = self._image_anchor_xy_emu(ws, img)
            if not pos:
                continue
            x_emu, _, row_idx = pos
            if header_rows[0] <= row_idx <= header_rows[1]:
                header_images.append((img, x_emu))

        to_remove: Set[Image] = set()
        if header_images:
            header_images.sort(key=lambda item: item[1])
            to_remove.add(header_images[0][0])
            to_remove.add(header_images[-1][0])

        ws._images = [img for img in getattr(ws, "_images", []) if img not in to_remove]

        emu_per_px = 9525
        margin_left_emu = int(round(18 * emu_per_px))
        margin_right_emu = int(round(24 * emu_per_px))
        gap_title_emu = int(round(12 * emu_per_px))

        # BIT (izquierda)
        bit_cx = int(ref_bit.get("cx", 0))
        bit_cy = int(ref_bit.get("cy", 0))
        if bit_cx > 0 and bit_cy > 0:
            bit_y = max(int(round((header_height_emu - bit_cy) / 2)), 0)
            bit_x = margin_left_emu
            max_left = max(title_x_start_emu - gap_title_emu, 0)
            x_max = max_left - bit_cx
            if bit_x > x_max:
                bit_x = max(0, x_max)
            if bit_x + bit_cx > max_left:
                available = max_left
                scale = min(
                    available / float(bit_cx) if bit_cx else 0.0,
                    header_height_emu / float(bit_cy) if bit_cy else 0.0,
                ) * 0.98
                if scale > 0 and scale < 1:
                    bit_cx = max(int(bit_cx * scale), 1)
                    bit_cy = max(int(bit_cy * scale), 1)
                    bit_y = max(int(round((header_height_emu - bit_cy) / 2)), 0)
                    bit_x = max(0, max_left - bit_cx)
                    logger.info("Escalado de emergencia aplicado al logo BIT (%.2f).", scale)
            self.place_logo_fixed(ws, bit_logo_path, bit_cx, bit_cy, bit_x, bit_y, header_rows[0])

        # NOVA (derecha)
        nova_cx = int(ref_nova.get("cx", 0))
        nova_cy = int(ref_nova.get("cy", 0))
        if nova_cx > 0 and nova_cy > 0:
            nova_y = max(int(round((header_height_emu - nova_cy) / 2)), 0)
            nova_x = header_width_emu - margin_right_emu - nova_cx
            min_right = title_x_end_emu + gap_title_emu
            if nova_x < min_right:
                nova_x = min_right
            if nova_x + nova_cx > header_width_emu:
                nova_x = max(0, header_width_emu - nova_cx)
            if nova_x < min_right:
                available = max(header_width_emu - min_right, 0)
                scale = min(
                    available / float(nova_cx) if nova_cx else 0.0,
                    header_height_emu / float(nova_cy) if nova_cy else 0.0,
                ) * 0.98
                if scale > 0 and scale < 1:
                    nova_cx = max(int(nova_cx * scale), 1)
                    nova_cy = max(int(nova_cy * scale), 1)
                    nova_y = max(int(round((header_height_emu - nova_cy) / 2)), 0)
                    nova_x = max(0, header_width_emu - margin_right_emu - nova_cx)
                    logger.info("Escalado de emergencia aplicado al logo NOVA (%.2f).", scale)
            self.place_logo_fixed(ws, nova_logo_path, nova_cx, nova_cy, nova_x, nova_y, header_rows[0])

    def _get_image_dimensions_px(self, img_path: Path) -> Optional[Tuple[int, int]]:
        if not img_path or not img_path.exists():
            return None
        try:
            if PILImage:
                with PILImage.open(img_path) as pil_img:
                    return (pil_img.width, pil_img.height)
            img = Image(str(img_path))
            if img.width and img.height:
                return (int(img.width), int(img.height))
        except Exception as exc:
            logger.warning("No se pudo leer el tamaño del logo %s: %s", img_path, exc)
        return None

    def _compute_target_size(
        self,
        img_path: Path,
        max_width: int,
        max_height: int,
    ) -> Optional[Tuple[int, int]]:
        dims = self._get_image_dimensions_px(img_path)
        if not dims:
            return None
        img_w, img_h = dims
        if img_w <= 0 or img_h <= 0:
            return None
        width_scale = max_width / float(img_w) if max_width else 1.0
        height_scale = max_height / float(img_h) if max_height else width_scale
        scale = min(width_scale, height_scale)
        return (int(img_w * scale), int(img_h * scale))

    def get_header_metrics(
        self,
        ws: Worksheet,
        title_text: str = "Informe de Actividades",
    ) -> Dict[str, object]:
        """Calcula métricas del header para posicionar logos con márgenes."""
        title_merge = self._find_title_merge(ws, title_text)
        if title_merge:
            title_min_row, title_max_row, title_min_col, title_max_col = title_merge
            header_start = 1
            header_end = title_max_row
            max_scan = min(title_max_row + 6, ws.max_row)
            for r in range(title_max_row + 1, max_scan + 1):
                has_fill = False
                for c in range(title_min_col, title_max_col + 1):
                    cell = ws.cell(row=r, column=c)
                    if cell.fill and cell.fill.fill_type:
                        has_fill = True
                        break
                if has_fill:
                    header_end = r
                else:
                    break
        else:
            header_start = 1
            header_end = min(3, ws.max_row or 3)
            title_min_row = header_start
            title_max_row = header_end
            title_min_col = 1
            title_max_col = 1

        header_rows = (header_start, header_end)
        last_col = self._get_last_header_col(ws, header_rows, title_merge)

        header_height_px = 0
        default_height = getattr(ws.sheet_format, "defaultRowHeight", 15.0)
        for r in range(header_start, header_end + 1):
            header_height_px += self._row_height_to_pixels(ws.row_dimensions[r].height, default_height)

        header_width_px = 0
        for c in range(1, last_col + 1):
            col_letter = get_column_letter(c)
            header_width_px += self._column_width_to_pixels(ws.column_dimensions[col_letter].width)

        title_x_start_px = 0
        title_x_end_px = header_width_px
        if title_merge:
            title_x_start_px = 0
            for c in range(1, title_min_col):
                col_letter = get_column_letter(c)
                title_x_start_px += self._column_width_to_pixels(ws.column_dimensions[col_letter].width)
            title_x_end_px = title_x_start_px
            for c in range(title_min_col, title_max_col + 1):
                col_letter = get_column_letter(c)
                title_x_end_px += self._column_width_to_pixels(ws.column_dimensions[col_letter].width)

        if header_width_px > 0:
            # Cuando el título ocupa casi todo el ancho, usa una zona central segura
            if title_x_start_px <= header_width_px * 0.1 and title_x_end_px >= header_width_px * 0.9:
                title_x_start_px = int(round(header_width_px * 0.35))
                title_x_end_px = int(round(header_width_px * 0.65))

        return {
            "headerRows": header_rows,
            "headerHeightPx": header_height_px,
            "headerWidthPx": header_width_px,
            "titleXStartPx": title_x_start_px,
            "titleXEndPx": title_x_end_px,
            "lastColUsedInHeader": last_col,
        }

    def place_logo_fixed(
        self,
        ws: Worksheet,
        img_path: Path,
        cx: int,
        cy: int,
        x_emu: int,
        y_emu: int,
        header_start_row: int,
    ) -> None:
        """Inserta un logo en posición absoluta (EMU) respecto al inicio del header."""
        if not img_path or not img_path.exists():
            return
        if cx <= 0 or cy <= 0:
            return

        img = Image(str(img_path))
        col_idx0, row_idx0, col_off, row_off = self.xy_to_anchor(ws, x_emu, y_emu, start_row=header_start_row)
        marker = AnchorMarker(
            col=col_idx0,
            colOff=col_off,
            row=row_idx0,
            rowOff=row_off,
        )
        ext = XDRPositiveSize2D(int(cx), int(cy))
        img.anchor = OneCellAnchor(_from=marker, ext=ext)
        ws.add_image(img)
