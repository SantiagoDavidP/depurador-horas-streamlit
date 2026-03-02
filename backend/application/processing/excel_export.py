from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Optional

import pandas as pd
from openpyxl.styles import Font, PatternFill

from backend.application.processing.constants import _CRITICAL_ERROR_TYPES
from .models import ProcessorSummary


class ProcessingExcelExportMixin:
    def _create_executive_summary_sheet(
        self,
        workbook,
        summary: ProcessorSummary,
    ) -> None:
        if workbook.sheetnames:
            ws = workbook[workbook.sheetnames[0]]
            ws.title = "Resumen Ejecutivo"
            ws.delete_rows(1, ws.max_row)
        else:
            ws = workbook.create_sheet("Resumen Ejecutivo", 0)

        header_font = Font(size=14, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        success_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
        danger_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        ws["A1"] = "Reporte de Validación de Timesheet"
        ws["A1"].font = Font(size=18, bold=True, color="1E3A8A")
        ws.merge_cells("A1:E1")

        ws["A2"] = f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        ws.merge_cells("A2:E2")

        row = 4
        ws[f"A{row}"] = "Métricas clave"
        ws[f"A{row}"].font = header_font
        ws[f"A{row}"].fill = header_fill
        ws.merge_cells(f"A{row}:E{row}")
        row += 1

        metrics = [
            (
                "Score de Calidad",
                f"{summary.quality_score:.0f}/100",
                "🟢 Excelente" if summary.quality_score >= 90 else
                "🟡 Bueno" if summary.quality_score >= 70 else
                "🔴 Atención",
            ),
            ("Registros procesados", summary.total_registros, f"{summary.metadata_removidas} metadata removidas"),
            ("Horas totales", f"{summary.horas_totales:.1f}h", f"{summary.horas_totales / 8:.1f} días"),
            ("Errores críticos", summary.errores_criticos, "Bloquean aprobación" if summary.errores_criticos else "Ninguno"),
        ]

        for label, value, detail in metrics:
            ws[f"A{row}"] = label
            ws[f"A{row}"].font = Font(bold=True)
            ws[f"B{row}"] = value
            ws[f"C{row}"] = detail
            row += 1

        row += 1
        ws[f"A{row}"] = "Distribución de errores"
        ws[f"A{row}"].font = header_font
        ws[f"A{row}"].fill = header_fill
        ws.merge_cells(f"A{row}:C{row}")
        row += 1

        ws[f"A{row}"] = "Tipo de error"
        ws[f"B{row}"] = "Cantidad"
        ws[f"C{row}"] = "Severidad"
        for col in ("A", "B", "C"):
            ws[f"{col}{row}"].font = Font(bold=True)
            ws[f"{col}{row}"].fill = PatternFill(start_color="E5E7EB", end_color="E5E7EB", fill_type="solid")
        row += 1

        if summary.errores_por_tipo:
            for tipo, cantidad in summary.errores_por_tipo.items():
                ws[f"A{row}"] = tipo.replace("_", " ").title()
                ws[f"B{row}"] = cantidad
                ws[f"C{row}"] = "Crítico" if tipo in _CRITICAL_ERROR_TYPES else "Advertencia"
                row += 1
        else:
            ws[f"A{row}"] = "Sin errores registrados"
            ws.merge_cells(f"A{row}:C{row}")
            row += 1

        row += 1
        ws[f"A{row}"] = "Decisión"
        ws[f"A{row}"].font = header_font
        ws[f"A{row}"].fill = header_fill
        ws.merge_cells(f"A{row}:C{row}")
        row += 1

        ws[f"A{row}"] = (
            "✅ Timesheet aprobable" if summary.errores_criticos == 0 else
            f"🚫 Requiere {summary.errores_criticos} correcciones críticas"
        )
        ws[f"A{row}"].font = Font(size=14, bold=True, color="047857" if summary.errores_criticos == 0 else "991B1B")
        ws[f"A{row}"].fill = success_fill if summary.errores_criticos == 0 else danger_fill
        ws.merge_cells(f"A{row}:C{row}")
        row += 2

        if summary.ai_summary:
            ws[f"A{row}"] = "Análisis inteligente"
            ws[f"A{row}"].font = header_font
            ws[f"A{row}"].fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
            ws.merge_cells(f"A{row}:C{row}")
            row += 1

            ws[f"A{row}"] = "Diagnóstico:"
            ws[f"A{row}"].font = Font(bold=True)
            ws[f"B{row}"] = summary.ai_summary.get("diagnostico", "")
            ws.merge_cells(f"B{row}:C{row}")
            row += 1

            acciones = summary.ai_summary.get("acciones") or []
            if acciones:
                ws[f"A{row}"] = "Acciones recomendadas:"
                ws[f"A{row}"].font = Font(bold=True)
                row += 1
                for idx, accion in enumerate(acciones, 1):
                    ws[f"A{row}"] = f"{idx}. {accion}"
                    ws.merge_cells(f"A{row}:C{row}")
                    row += 1

            tiempo = summary.ai_summary.get("tiempo_estimado")
            if tiempo:
                ws[f"A{row}"] = "Tiempo estimado:"
                ws[f"A{row}"].font = Font(bold=True)
                ws[f"B{row}"] = tiempo
                ws.merge_cells(f"B{row}:C{row}")

        ws.column_dimensions["A"].width = 40
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 30

    def _export_workbook(
        self,
        corrected_df: pd.DataFrame,
        errors_df: pd.DataFrame,
        summary: ProcessorSummary,
        debug_hours_df: Optional[pd.DataFrame] = None,
        debug_detail_df: Optional[pd.DataFrame] = None,
    ) -> bytes:
        def _sanitize_excel_df(df: pd.DataFrame) -> pd.DataFrame:
            if df is None or df.empty:
                return df
            return df.where(pd.notna(df), None)

        corrected_df = _sanitize_excel_df(corrected_df)
        errors_df = _sanitize_excel_df(errors_df)
        if debug_hours_df is not None:
            debug_hours_df = _sanitize_excel_df(debug_hours_df)
        if debug_detail_df is not None:
            debug_detail_df = _sanitize_excel_df(debug_detail_df)

        buffer = BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            workbook = writer.book
            self._create_executive_summary_sheet(workbook, summary)
            corrected_df.to_excel(writer, index=False, sheet_name="Datos Corregidos")
            errors_df.to_excel(writer, index=False, sheet_name="Reporte de Errores")
            if debug_hours_df is not None and not debug_hours_df.empty:
                debug_hours_df.to_excel(writer, index=False, sheet_name="Debug Horas")
            if debug_detail_df is not None and not debug_detail_df.empty:
                debug_detail_df.to_excel(writer, index=False, sheet_name="Debug Filas")
        buffer.seek(0)
        return buffer.getvalue()

